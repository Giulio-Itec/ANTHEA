from zipfile import ZipFile
from lxml import etree
from docx import Document
import json
p=r'C:/Users/g.pacini/Desktop/MODELLO-RELAZIONE-ITEC-AA.docx'
d=Document(p)
for i,s in enumerate(d.sections): print('SECTION',i, s.page_width.mm,s.page_height.mm, 'margins',s.top_margin.mm,s.bottom_margin.mm,s.left_margin.mm,s.right_margin.mm)
for i,p in enumerate(d.paragraphs):
 if p.text: print('P',i,p.style.name,repr(p.text))
for i,t in enumerate(d.tables):print('TABLE',i, [[c.text for c in r.cells] for r in t.rows])
for s in d.styles:
 if s.type==1 and (s.name in ['Normal','Title','Subtitle'] or 'heading' in s.name.lower()):print('STYLE',s.name,s.style_id,s.font.name,s.font.size.pt if s.font.size else None)
with ZipFile(p if isinstance(p,str) else r'C:/Users/g.pacini/Desktop/MODELLO-RELAZIONE-ITEC-AA.docx') as z:
 for n in z.namelist():
  if n.startswith(('word/header','word/footer')) and n.endswith('.xml'): print(n, etree.fromstring(z.read(n)).xpath('//*[local-name()="t"]/text()'))
