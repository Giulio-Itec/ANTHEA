"""Freeze independently acquired Excel values and provenance. Refuses to replace existing evidence."""
import sys,json,hashlib,gzip,shutil
from pathlib import Path
import openpyxl
root,checker=map(Path,sys.argv[1:]); source=root/'supporto/artefatti/spettri-ntc'; out=checker/'tests/Design.Checks/Fixtures'
out.mkdir(parents=True,exist_ok=True)
def save(name,value):
    path=out/name
    if path.exists():raise RuntimeError(f'Atteso già esistente: {path}')
    path.write_text(json.dumps(value,ensure_ascii=False,indent=2),encoding='utf-8')
capture=json.loads((source/'excel-capture-01.json').read_text(encoding='utf-8-sig'))
save('spettri-cslp103-excel.json',capture)
book=source/'fonti/isolamento-ntc18.xlsx';w=openpyxl.load_workbook(book,read_only=True,data_only=True)
def cells(sheet,addresses):return {a:w[sheet][a].value for a in addresses}
save('spettri-isolamento-excel.json',{
    'sha256':hashlib.sha256(book.read_bytes()).hexdigest(),
    'isolation':cells('Isolatori',['H5','H6','H7','H8','H9','H10','H11','H12','H13','H26','H27','H28','H29','H30','H33','H34','H35','H36','H37']),
    'hazard':[[w['Isolatori'].cell(r,c).value for c in [3,4,5]] for r in range(18,22)],
    'pendulum':cells('Modellazione_isolatori_pendolo',['H4','H5','H6','H8','H9','H11','H13','H17']),
    'spectra':{s:[[w[s].cell(r,c).value for c in [7,8]] for r in range(4,49)] for s in ['SLO','SLD','SLV','SLC']}})
files=[source/'fonti/cslp/Spettri-NTCver.1.0.3.xls',source/'cslp-xlsx/reference.xlsx',book]
save('spettri-sources.json',{'official_url':'https://cslp.mit.gov.it/pareri/spettri-ntc-ver-103','capture':'Excel 16 COM; macros disabled; analysis copy; Capture-Reference.ps1','files':[{ 'name':p.name,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in files]})
