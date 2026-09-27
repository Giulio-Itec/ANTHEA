"""Build the two versioned user manuals from their maintainable Markdown sources.

Run with the Codex bundled Python. Rendering/inspection is a separate QA step.
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'supporto/documentazione/Guide_ANTHEA'
OUT.mkdir(parents=True, exist_ok=True)

def field(paragraph, code):
    r=OxmlElement('w:r'); f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'), code)
    t=OxmlElement('w:t'); t.text='1'; r.append(t); f.append(r); paragraph._p.append(f)

def bookmark(paragraph, name, index):
    start=OxmlElement('w:bookmarkStart'); start.set(qn('w:id'),str(index)); start.set(qn('w:name'),name)
    end=OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'),str(index))
    paragraph._p.insert(0,start); paragraph._p.append(end)

def link(paragraph, text, anchor):
    h=OxmlElement('w:hyperlink'); h.set(qn('w:anchor'),anchor)
    r=OxmlElement('w:r'); p=OxmlElement('w:rPr'); c=OxmlElement('w:color'); c.set(qn('w:val'),'173858'); p.append(c)
    r.append(p); t=OxmlElement('w:t'); t.text=text; r.append(t); h.append(r); paragraph._p.append(h)

def write_table(doc, lines):
    rows=[[c.strip() for c in line.strip().strip('|').split('|')] for line in lines]
    rows=[r for r in rows if not all(re.fullmatch(r'[:\- ]+',c) for c in r)]
    table=doc.add_table(rows=0,cols=len(rows[0])); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
    n=len(rows[0]); widths=([5,12] if n==2 else [4,6,7] if n==3 else [4.7,3,3,3.3,3] if n==5 else [17/n]*n)
    for col,w in zip(table.columns,widths): col.width=Cm(w)
    props=table._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
    props.append(borders)
    for idx, values in enumerate(rows):
        row=table.add_row()
        if idx==0:
            repeat=OxmlElement('w:tblHeader'); row._tr.get_or_add_trPr().append(repeat)
        no_split=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
        for j,(cell,value) in enumerate(zip(row.cells,values)):
            cell.width=Cm(widths[j]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr=cell._tc.get_or_add_tcPr(); shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'173858' if idx==0 else ('F2F5F8' if idx%2==0 else 'FFFFFF')); tcPr.append(shade)
            margins=OxmlElement('w:tcMar')
            for edge in ('top','left','bottom','right'):
                m=OxmlElement('w:'+edge); m.set(qn('w:w'),'100'); m.set(qn('w:type'),'dxa'); margins.append(m)
            tcPr.append(margins)
            p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(2); p.paragraph_format.space_before=Pt(2); p.paragraph_format.line_spacing=1.05
            if j>0 and len(value)<18: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            run=p.add_run(value); run.font.size=Pt(9); run.bold=idx==0; run.font.color.rgb=RGBColor.from_string('FFFFFF' if idx==0 else '000000')
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

def build(stem, filename, short):
    source=ROOT/'supporto/docs'/f'{stem}.md'; lines=source.read_text(encoding='utf-8').splitlines()
    doc=Document(); sec=doc.sections[0]
    sec.page_width=Cm(21); sec.page_height=Cm(29.7); sec.top_margin=Cm(1.9); sec.bottom_margin=Cm(1.9); sec.left_margin=Cm(2); sec.right_margin=Cm(2)
    sec.header_distance=Cm(.8); sec.footer_distance=Cm(.8)
    normal=doc.styles['Normal']; normal.font.name='Calibri'; normal.font.size=Pt(10.5)
    normal.paragraph_format.line_spacing=1.12; normal.paragraph_format.space_after=Pt(7)
    for name,size in [('Title',26),('Subtitle',12),('Heading 1',16),('Heading 2',12),('Heading 3',11)]:
        s=doc.styles[name]; s.font.name='Calibri'; s.font.size=Pt(size); s.font.color.rgb=RGBColor(0,0,0)
        s.font.bold=name.startswith('Heading'); s.paragraph_format.keep_with_next=True
        s.paragraph_format.space_before=Pt(15 if name.startswith('Heading') else 3); s.paragraph_format.space_after=Pt(8)
    # The bundled default template carries a blue bottom border on Title.
    for border in doc.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    h=sec.header.paragraphs[0]; h.add_run('ANTHEA  ·  '+short).font.size=Pt(8); h.paragraph_format.space_after=Pt(0)
    f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    f.add_run('Edizione 1 · 26 settembre 2026     |     ').font.size=Pt(8); field(f,'PAGE')
    settings=doc.settings.element; update=OxmlElement('w:updateFields'); update.set(qn('w:val'),'true'); settings.append(update)
    doc.core_properties.title=lines[0][2:]; doc.core_properties.subject=short; doc.core_properties.author='ANTHEA'; doc.core_properties.keywords='Manuale; ANTHEA; 2026-09-26'
    headings=[l[3:] for l in lines if l.startswith('## ')]
    i=0; chapter=0; first=True
    while i<len(lines):
        line=lines[i].strip(); i+=1
        if not line: continue
        if line.startswith('# '): doc.add_paragraph(line[2:],'Title'); continue
        if line.startswith('## '):
            if first:
                doc.add_heading('Percorso di lettura',2)
                for k,title in enumerate(headings):
                    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(3); link(p,title,f'chapter_{k}')
                doc.add_page_break(); first=False
            p=doc.add_heading(line[3:],1); bookmark(p,f'chapter_{chapter}',chapter+1); chapter+=1; continue
        if line.startswith('### '): doc.add_heading(line[4:],2); continue
        if line.startswith('|'):
            block=[line]
            while i<len(lines) and lines[i].strip().startswith('|'): block.append(lines[i]); i+=1
            write_table(doc,block); continue
        if line.startswith('$$ '):
            p=doc.add_paragraph(); p.paragraph_format.left_indent=Cm(.45); p.paragraph_format.space_after=Pt(3); p.paragraph_format.space_before=Pt(3)
            p.paragraph_format.keep_with_next=i<len(lines) and lines[i].startswith('$$ ')
            # Editable Word math, with symbols preserved and no raster text.
            math=OxmlElement('m:oMath'); r=OxmlElement('m:r'); prop=OxmlElement('m:rPr'); sty=OxmlElement('m:sty'); sty.set(qn('m:val'),'p'); prop.append(sty); r.append(prop)
            wr=OxmlElement('w:rPr'); sz=OxmlElement('w:sz'); sz.set(qn('w:val'),'20'); wr.append(sz); r.append(wr)
            t=OxmlElement('m:t'); t.set(qn('xml:space'),'preserve'); t.text=line[3:]; r.append(t); math.append(r); p._p.append(math)
            continue
        if line.startswith('!['):
            match=re.fullmatch(r'!\[(.*?)\]\((.*?)\)',line)
            path=(source.parent/match.group(2)).resolve()
            if not path.is_file(): raise FileNotFoundError(path)
            p=doc.add_paragraph(); p.paragraph_format.keep_with_next=True; pic=p.add_run().add_picture(str(path),width=Cm(17)); pic._inline.docPr.set('descr',match.group(1))
            p=doc.add_paragraph(match.group(1),'Caption'); p.style.font.color.rgb=RGBColor(0,0,0); continue
        p=doc.add_paragraph()
        if re.match(r'^\d+\. ',line): p.paragraph_format.left_indent=Cm(.5); p.paragraph_format.first_line_indent=Cm(-.5)
        p.add_run(line)
    target=OUT/filename; doc.save(target)
    print(f'{target}: {len(doc.paragraphs)} paragrafi, {len(doc.tables)} tabelle, {len(source.read_text(encoding="utf-8").split())} parole circa')

if __name__=='__main__':
    build('guida-pratica-anthea','ANTHEA_Guida_pratica_Rev01.docx','Guida pratica')
    build('guida-teorica-anthea','ANTHEA_Guida_teorica_Rev01.docx','Guida teorica')
