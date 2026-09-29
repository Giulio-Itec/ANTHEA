"""Rebuild the factual QA report from the saved ten-case manifest, without computing MAX results."""
from pathlib import Path
import sys, json, hashlib
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

root = Path(sys.argv[1]).resolve()
cases = json.loads((root / 'indice-confronti.json').read_text(encoding='utf-8-sig'))
max_status = json.loads((root / 'stato-max.json').read_text(encoding='utf-8-sig')) if (root / 'stato-max.json').exists() else {'completed': 0, 'note': 'Il confronto MAX non è documentato in questa raccolta.'}
assert len(cases) == 10
assert (root / 'stato-esecuzione.txt').read_text().startswith('PASS')
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.2677), Inches(11.6929)
sec.top_margin = sec.bottom_margin = Inches(.7)
sec.left_margin = sec.right_margin = Inches(.6)
for style in doc.styles:
    for node in list(style.element.iter(qn('w:pBdr'))): node.getparent().remove(node)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(4)
doc.styles['Title'].font.size = Pt(23)
doc.styles['Heading 1'].font.size = Pt(15)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(10)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(4)
doc.styles['Subtitle'].font.italic = False

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for cell, width in zip(t.columns, widths): cell.width = Inches(width)
    for i, text in enumerate(headers): t.rows[0].cells[i].text = text
    trpr = t.rows[0]._tr.get_or_add_trPr(); repeat = OxmlElement('w:tblHeader'); trpr.append(repeat)
    for row in rows:
        for cell, text in zip(t.add_row().cells, row): cell.text = str(text)
    for ri, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
        for i, cell in enumerate(row.cells):
            cell.width = Inches(widths[i])
            pr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for side in ['top', 'left', 'bottom', 'right']:
                edge = OxmlElement('w:' + side); edge.set(qn('w:val'), 'single'); edge.set(qn('w:sz'), '4'); edge.set(qn('w:color'), 'D9D9D9'); borders.append(edge)
            pr.append(borders)
            if ri == 0:
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E8EFF7'); pr.append(shade)
            margin = OxmlElement('w:tcMar')
            for side in ['top', 'bottom', 'left', 'right']:
                e = OxmlElement('w:' + side); e.set(qn('w:w'), '70'); e.set(qn('w:type'), 'dxa'); margin.append(e)
            pr.append(margin)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs: run.font.size = Pt(9.5); run.bold = ri == 0
    doc.add_paragraph()

doc.add_paragraph('Controllo della stabilità globale', 'Title')
doc.add_paragraph('Modulo muri di sostegno ANTHEA', 'Subtitle')
doc.add_paragraph('29 settembre 2026')
doc.add_paragraph('La stabilità globale è implementata con motore Bishop e geometria separati. Sono disponibili dieci modelli ANTHEA riproducibili, risultati numerici e relazioni. I controlli interni sono superati; il confronto indipendente con MAX 16 non è stato eseguito.')
doc.add_paragraph(max_status['note'] + ' Confronti MAX completati e documentati: ' + str(max_status['completed']) + '/10. Nessun valore calcolato da ANTHEA viene attribuito a MAX.')
doc.add_heading('Cosa è stato controllato', 1)
doc.add_paragraph('Il motore risolve le superfici circolari sotto l’intero muro, con stratigrafia orizzontale, pesi saturi, falda, azioni esterne e inerzie pseudostatiche. Il muro sostituisce il terreno nel proprio volume: spinte e reazioni interne non vengono sommate una seconda volta. La resistenza può essere espressa con parametri drenati oppure cu, entro i limiti numerici dichiarati.')
doc.add_paragraph('I controlli includono casi analitici a un concio, equilibrio verticale e dei momenti, area e baricentro GPC, peso del muro, coefficienti parziali, ricerca più densa, input incompleti, matrici obsolete, cancellazione, salvataggio e rilettura dei dieci archivi. I controlli software non equivalgono alla validazione indipendente con un altro programma.')
table(['Controllo', 'Esito'], [
    ['Numerica e dieci esempi', (root / 'stato-esecuzione.txt').read_text().split(';')[0]],
    ['Regressione modulo muri', '163 controlli superati'],
    ['Interfaccia stabilità globale', '14 controlli superati'],
    ['Regressione interfaccia muri', '41 controlli superati'],
    ['Dati progetti e materiali', '988 controlli superati'],
    ['Compilazione WPF', '0 errori e 0 avvisi'],
    ['Confronto MAX 16', 'Non eseguito'],
], [4.8, 2.2])

doc.add_heading('Risultati dei dieci esempi', 1)
doc.add_paragraph('F è calcolato con i parametri della combinazione; il tasso di lavoro è γR/F. Per i casi 08 e 09 la tabella mostra rispettivamente il caso sismico e quello eccezionale. Per gli altri casi mostra il maggiore tasso numerico. I risultati dettagliati di tutte le combinazioni si trovano nei file allegati.')
rows = []
for item in cases:
    relevant = item['Cases']
    if item['Id'].startswith('08'): relevant = [c for c in relevant if 'SLV' in c['Name']]
    if item['Id'].startswith('09'): relevant = [c for c in relevant if 'eccezionale' in c['Name']]
    worst = max(relevant, key=lambda c: c['Eta'] or 0)
    status = 'Incompleta' if worst['NumericalFailures'] or worst['Boundary'] else 'Non soddisfatta' if worst['Eta'] > 1 else 'Soddisfatta nel dominio'
    rows.append([item['Id'].replace('-', ' '), f"{worst['F']:.4f}".replace('.', ','), f"{worst['Eta']:.4f}".replace('.', ','), status, 'Da eseguire'])
table(['Esempio', 'F', 'γR/F', 'Esito ANTHEA', 'MAX 16'], rows, [2.1, .7, .7, 2.3, 1.2])
doc.add_paragraph('Il caso 10 non drenato non produce una verifica favorevole: alcuni conci non ammettono una soluzione senza trazione e il controllo della superficie minima con il doppio dei conci non conferma una soluzione ammissibile. Il tasso numerico riportato non supera questa limitazione. È un esempio di segnalazione esplicita di analisi incompleta.')
doc.add_heading('Modifiche emerse dal controllo', 1)
doc.add_paragraph('La prima ricerca raffinava soltanto il minimo della griglia iniziale. Nel caso base senza sovraccarico il risultato era F=1,337453, mentre una griglia diversa individuava F=1,254789. La ricerca è stata modificata per raffinare sei minimi distinti, con quattro livelli nel preset.')
doc.add_paragraph('Dopo la correzione, la griglia 9³ restituisce F=1,235742 e la griglia 13³ con più conci restituisce F=1,234707: differenza circa 0,084%. Questo controllo documenta la sensibilità del caso base; non garantisce il minimo assoluto per altri modelli.')
doc.add_paragraph('Sono stati separati i coefficienti della stabilità globale da quelli delle spinte. La scelta Wood non porta βs a 1: il preset globale SLV usa βs=0,38, mentre il modello delle spinte conserva le proprie regole. Il coefficiente γR entra una sola volta nel tasso di lavoro.')
doc.add_paragraph('Le matrici globali modificate sono conservate e rese non utilizzabili quando cambiano azioni, coefficienti ψ o dati sismici rilevanti. Gli input errati cancellano i risultati precedenti. Il calcolo globale può essere avviato anche quando i dati delle verifiche locali sono fuori campo.')
doc.add_heading('Organizzazione del codice e riuso GPC', 1)
table(['Componente', 'Posizione e destinazione proposta'], [
    ['Geometria', 'SlopeGeometry.cs; costruzioni e interrogazioni di profili. Candidato per GPC.Geometry. Area e baricentro usano già GPC.Geometry.'],
    ['Materiale geotecnico', 'SlopeSoil in SlopeStability.Models.cs. Candidato per GPC.Model; nei sorgenti disponibili non è presente un materiale con φ′, c′ e cu.'],
    ['Equilibrio', 'BishopSolver.cs; nessuna dipendenza da muro, JSON, interfaccia o normativa. Candidato per GPC di geotecnica.'],
    ['Ricerca e discretizzazione', 'SlopeStability.cs; separata dalla geometria e dall’adattatore del muro.'],
    ['Adattatore ANTHEA', 'RetainingWall.GlobalStability.cs; documento, azioni e preset NTC. GPC.Model associa azioni e coefficienti.'],
], [1.8, 5.2])
doc.add_paragraph('I materiali strutturali e le verifiche delle sezioni restano nel percorso GPC esistente. La globale utilizza il peso specifico del materiale del muro; non impiega moduli elastici fittizi per rappresentare i terreni. Non sono stati modificati i repository esterni GPC.')
doc.add_heading('Come ripetere i controlli', 1)
doc.add_paragraph('Ogni cartella esempi contiene modello.anthea, risultati-anthea.json, conci-anthea.csv, relazione-anthea.docx e riproduzione.md. Aprire il modello, consultare Input → Terreno → Stabilità globale, quindi Verifiche → Stabilità globale. Il modello conserva profilo, strati, falda, dominio e combinazioni.')
doc.add_paragraph('Per rigenerare la raccolta eseguire il progetto supporto/test/GlobalStability.Checks indicando una cartella di output. Il file controlli-numerici.txt elenca gli assert eseguiti; stato-esecuzione.txt distingue una corsa completa da una rimasta in corso.')
doc.add_heading('Confronto MAX ancora da completare', 1)
for text in [
    'Usare Bishop e uniformare unità, quote, condizioni drenate o non drenate, falda, carichi e pesi. Il muro deve essere incluso nella massa senza duplicare le spinte interne.',
    'Confrontare prima lo stesso cerchio usando centro e raggio salvati nel JSON. Allineare i parametri ridotti M2 oppure i parametri caratteristici: non confrontare direttamente fattori di sicurezza riferiti a set diversi.',
    'Confrontare poi la ricerca, documentando limiti del dominio, numero di conci, esclusioni geometriche e criteri di convergenza. Registrare la versione precisa di MAX 16.',
    'Salvare un file nativo MAX per ciascun caso, risultati e schermate. Riportare scostamento assoluto e percentuale, motivazione, eventuale modifica al motore e confronto ripetuto. Non fissare un esito positivo finché questa fase non è stata svolta.',
]: doc.add_paragraph(text, style='List Number')
doc.add_heading('Campo e riferimenti', 1)
doc.add_paragraph('La ricerca riguarda il ramo inferiore di cerchi con centro fra ingresso e uscita, base entro ±80° e passaggio sotto l’intero muro. Non comprende superfici non circolari, strati inclinati, fessure di trazione, acqua esterna, pressioni idrodinamiche, degradazione ciclica, liquefazione, cedimenti o stabilità autonoma del versante naturale. Un minimo sul bordo richiede l’ampliamento della ricerca.')
doc.add_paragraph('Preset: NTC 2018 §§6.5.3.1.1 e 6.8.2 per A2–M2–R2; §§7.11.4 e 7.11.6.2.2 per SLV del complesso muro–terreno. Fonte: https://www.gazzettaufficiale.it/eli/gu/2018/02/20/42/so/8/sg/pdf')
doc.add_paragraph('Metodo: Bishop A.W. (1955), The use of the slip circle in the stability analysis of slopes; USACE EM 1110-2-1902, Slope Stability, appendice C. Fonte: https://www.publications.usace.army.mil/Portals/76/Publications/EngineerManuals/EM_1110-2-1902.pdf')
footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('ANTHEA  |  Controllo interno  |  ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
out = root / 'Rapporto-controllo-stabilita-globale.docx'; doc.save(out)
hashes = {str(p.relative_to(root)): hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted((root / 'esempi').rglob('*')) if p.is_file()}
(root / 'impronte-sha256.json').write_text(json.dumps(hashes, indent=2), encoding='utf-8')
print(out)
