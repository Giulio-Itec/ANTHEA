"""Build the bridge calculation validation dossier from executed test evidence.
Run with the Codex dependency Python. Charts use task-local Matplotlib.
No expected numerical result is generated from production code here.
"""
from pathlib import Path
import sys, json, math, re, hashlib, xml.etree.ElementTree as ET
from collections import Counter

ROOT = Path(__file__).resolve().parents[2]
ART = ROOT / "supporto/artefatti/ponte_curve_validazione"
sys.path.insert(0, str(ART / "plot-runtime"))
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

OUT = ROOT / "supporto/documentazione/Validazione_Sezione_Ponte"
OUT.mkdir(parents=True, exist_ok=True)
FIG = ART / "figure"; FIG.mkdir(parents=True, exist_ok=True)
CHECKER = ROOT.parent / "Checker"
VAL = CHECKER / "GPCChecker.Test.CompositeBridge/Validation"
def load(p): return json.loads(Path(p).read_text(encoding="utf-8-sig"))
def fmt(x, digits=6):
    if x == 0: return "0"
    return f"{x:.{digits}g}".replace(".", ",")
def sha(path): return hashlib.sha256(Path(path).read_bytes()).hexdigest().upper()
NS = {"t": "http://microsoft.com/schemas/VisualStudio/TeamTest/2010"}
def trx(name):
    t = ET.parse(ART / name)
    c = t.find(".//t:Counters", NS).attrib
    rows = t.findall(".//t:UnitTestResult", NS)
    definitions = {r.attrib["id"]: r.find("t:TestMethod", NS).attrib["className"].split(",")[0].split(".")[-1]
                   for r in t.findall(".//t:UnitTest", NS)}
    classes = Counter(definitions[r.attrib["testId"]] for r in rows)
    return c, rows, classes
composite, cr, cc = trx("composite-finale.trx")
audit, ar, ac = trx("audit-finale.trx")
bugs, br, bc = trx("difetti-noti.trx")
known_skipped = sum(r.attrib.get("outcome") == "NotExecuted" for r in br)
assert int(composite["failed"]) == 0 and int(audit["failed"]) == 0
total = int(composite["passed"]) + int(audit["passed"])
numbers = ART / "numeri"
cum = load(numbers / "dossier-cumulativo.json")
linear = load(numbers / "dossier-storico-lineare.json")
analytical = load(numbers / "dossier-analytical.json")
published = load(numbers / "dossier-published.json")
curves = load(VAL / "opensees-curves-3.8.0.json")["cases"]
actual = [load(numbers / f"curve-{i}.json") for i in range(len(curves))]
history = load(VAL / "opensees-3.8.0.json")["cases"]
history_errors = {}
for r in cr:
    out = r.find("t:Output/t:StdOut", NS)
    for name, stress, strain in re.findall(r"(\w+)\[\d+\]: maxDeltaStress=([\d.Ee+\-]+) MPa; maxDeltaStrain=([\d.Ee+\-]+)", out.text or "" if out is not None else ""):
        old = history_errors.setdefault(name, [0., 0.])
        old[0] = max(old[0], float(stress)); old[1] = max(old[1], float(strain))
assert len(history_errors) == 7
max_hist = max(v[0] for v in history_errors.values())
max_curve = max(c["maxStress"] for c in actual)

plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10, "axes.spines.top": False, "axes.spines.right": False,
                     "axes.labelcolor": "#182D43", "text.color": "#182D43", "axes.grid": True, "grid.alpha": .2})
def curve_plot(index, name, axial=False):
    ref, act = curves[index]["points"], actual[index]["points"]
    fig, ax = plt.subplots(figsize=(7, 3.6), layout="constrained")
    xkey, ykey = ("axial", "n") if axial else ("curvature", "mref")
    xf, yf = (1e6, .001) if axial else (1000, 1e-6)
    x = [p[xkey]*xf for p in ref]; y = [p[ykey]*yf for p in ref]
    ax.plot(x, y, "-o", color="#2C718A", lw=1.5, ms=5, label="OpenSees")
    ax.plot([p[xkey]*xf for p in act], [p[ykey]*yf for p in act], "x", color="#AD5B26", ms=8, mew=1.5, label="Checker")
    for i, (a,b) in enumerate(zip(x,y)): ax.annotate(str(i+1), (a,b), xytext=(6,5), textcoords="offset points", fontsize=8)
    ax.axhline(0, color="#8D9AA7", lw=.7); ax.axvline(0, color="#8D9AA7", lw=.7)
    ax.set_xlabel("Deformazione assiale ε₀ [µε]" if axial else "Curvatura κ [1/m]")
    ax.set_ylabel("N [kN]" if axial else "M a y = 0 [kNm]"); ax.legend(frameon=False)
    fig.savefig(FIG/name, dpi=220); plt.close(fig)
curve_plot(1, "momento_curvatura.png")
curve_plot(4, "forza_deformazione.png", True)
curve_plot(3, "composta_curvatura.png")
fig, ax = plt.subplots(figsize=(7,3.1), layout="constrained")
names = ["G1","G2","R1","Q","R2"]
for label, quantity, color in [("Acciaio inferiore","Acciaio · intradosso","#245E81"),("CLS superiore","Soletta · estradosso","#A35320")]:
    rows = cum["Rows"]
    y = [next((r["Actual"] for r in rows if r["Id"]==f"CU{i+1}" and r["Quantity"]==quantity), 0) for i in range(5)]
    ax.plot(names, y, "-o", label=label, color=color)
ax.set_ylabel("Tensione [MPa]"); ax.legend(frameon=False); fig.savefig(FIG/"fasi.png", dpi=220); plt.close(fig)

doc = Document()
sec = doc.sections[0]; sec.page_height=Cm(27.94); sec.page_width=Cm(21.59)
sec.top_margin=Cm(1.8); sec.bottom_margin=Cm(1.7); sec.left_margin=Cm(2); sec.right_margin=Cm(2)
sec.header_distance=Cm(.7); sec.footer_distance=Cm(.7)
styles=doc.styles
for style in styles:
    for border in list(style.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
for name in ["Normal","Title","Subtitle","Heading 1","Heading 2","Heading 3","Caption","Header","Footer"]:
    styles[name].font.name="Calibri"; styles[name].font.color.rgb=RGBColor(0,0,0)
    styles[name].font.underline=False
styles["Normal"].font.size=Pt(10.5)
styles["Normal"].paragraph_format.space_after=Pt(7); styles["Normal"].paragraph_format.line_spacing=1.08
styles["Title"].font.size=Pt(23); styles["Title"].font.bold=True
styles["Heading 1"].font.size=Pt(17); styles["Heading 1"].paragraph_format.space_after=Pt(10)
styles["Heading 2"].font.size=Pt(12); styles["Heading 2"].paragraph_format.space_before=Pt(12)
styles["Caption"].font.size=Pt(9); styles["Caption"].font.italic=False
header=sec.header.paragraphs[0]; header.text="ANTHEA   |   Sezione composta da ponte   |   Validazione numerica"
header.style="Header"; header.runs[0].font.size=Pt(8)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run("Revisione 01 · 25 settembre 2026   |   ")
field=OxmlElement("w:fldSimple"); field.set(qn("w:instr"),"PAGE"); footer._p.append(field)
for run in footer.runs: run.font.size=Pt(8)
doc.core_properties.title="Validazione dei metodi di calcolo della sezione composta da ponte"
doc.core_properties.subject="Cumulativo storico lineare storico non lineare curve momento curvatura e forza deformazione"
doc.core_properties.author="ANTHEA · verifica software assistita da Codex"

def p(text, style=None): return doc.add_paragraph(text, style)
def h(text): doc.add_heading(text, level=2)
def page(title): doc.add_page_break(); doc.add_heading(title, level=1)
def table(headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    borders=OxmlElement("w:tblBorders")
    for side in ["top","left","bottom","right","insideH","insideV"]:
        edge=OxmlElement("w:"+side); edge.set(qn("w:val"),"single")
        edge.set(qn("w:sz"),"4"); edge.set(qn("w:color"),"D9D9D9"); borders.append(edge)
    t._tbl.tblPr.append(borders)
    widths=widths or [17/len(headers)]*len(headers)
    for c, w in zip(t.columns,widths): c.width=Cm(w)
    for i, text in enumerate(headers): t.rows[0].cells[i].text=str(text)
    for row in rows:
        cells=t.add_row().cells
        for i, value in enumerate(row): cells[i].text=str(value)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr(); cant=OxmlElement("w:cantSplit"); trpr.append(cant)
        if ri==0:
            repeat=OxmlElement("w:tblHeader"); trpr.append(repeat)
        for ci,cell in enumerate(row.cells):
            cell.width=Cm(widths[ci]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props=cell._tc.get_or_add_tcPr(); shade=OxmlElement("w:shd");shade.set(qn("w:fill"),"E5EBF1" if ri==0 else ("F5F7F9" if ri%2==0 else "FFFFFF"));props.append(shade)
            margins=OxmlElement("w:tcMar")
            for side in ["top","left","bottom","right"]:
                node=OxmlElement("w:"+side);node.set(qn("w:w"),"75");node.set(qn("w:type"),"dxa");margins.append(node)
            props.append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(2); para.paragraph_format.space_before=Pt(2);para.paragraph_format.line_spacing=1
                for run in para.runs: run.font.size=Pt(9.3);run.font.bold=ri==0;run.font.color.rgb=RGBColor(0,0,0)
    p("")
    return t
def picture(path, caption, width=16.6):
    para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER;para.paragraph_format.keep_with_next=True
    para.add_run().add_picture(str(path), width=Cm(width))
    p(caption, "Caption")
def code(text):
    para=p(text)
    for r in para.runs: r.font.name="Consolas";r.font.size=Pt(8)
def metric_rows(rows):
    return [[r["Quantity"], fmt(r["Expected"]),fmt(r["Actual"]),fmt(abs(r["Actual"]-r["Expected"])),fmt(r["Tolerance"])] for r in rows]

# 1
doc.add_heading("Validazione dei metodi di calcolo della sezione composta da ponte",0)
p("Revisione 01   ·   25 settembre 2026", "Subtitle")
p(f"I benchmark definiti per i cinque metodi del modulo ponte hanno esito positivo: {total} casi delle suite ordinarie sono superati. Lo storico e le nuove curve sono confrontati anche con OpenSees 3.8.0, su 64 stati complessivi. L’evidenza riguarda i casi e le ipotesi descritti in questo documento; non attesta la validità generale del modello per qualsiasi ponte.")
p("Il documento è destinato al responsabile dello sviluppo e al progettista che deve valutare l’impiego del modulo. Raccoglie modello matematico, dati dei casi, risultati attesi e ottenuti, criteri di confronto, riproducibilità e difetti ancora aperti. Il perimetro è il modulo della sezione composta da ponte, comprese le sue verifiche accessorie.")
table(["Metodo","Evidenza principale","Esito e campo"],[
["Cumulativo precedente","Soluzioni analitiche e regressioni","Coerente con il modello a sezione efficace comune"],
["Storico lineare","Soluzioni analitiche e OpenSees elastico","Memoria al getto e ritiri verificati nei casi provati"],
["Storico non lineare","Soluzioni analitiche e OpenSees","Sezione lorda e legami dichiarati"],
["Momento curvatura","Analitica e 4 curve OpenSees","N costante e curvatura imposta"],
["Forza deformazione","Analitica e 2 curve OpenSees","Curvatura costante e deformazione imposta"]],[4.5,5.5,7])
h("Conclusione operativa")
p("I confronti indipendenti rafforzano la verifica dell’implementazione numerica. I modelli OpenSees e le scelte di discretizzazione sono stati predisposti nello stesso sviluppo: non sono una revisione eseguita da un progettista terzo né una validazione sperimentale.")
p(f"I casi dei difetti noti delle vecchie DLL sono eseguiti separatamente: {bugs['failed']} non superati e {known_skipped} ignorati. Non sono conteggiati come successi. Le curve non lineari non comprendono l’instabilità locale di classe 4, la viscosità nel tempo o una verifica SLU completa.")

# 2
page("Modelli e convenzioni comuni")
p("Il calcolo numerico risiede in GPCChecker.CompositeBridge. ANTHEA gestisce ingressi, selezione del metodo, risultati, grafici e archivi. Le leggi e i dati di catalogo provengono da Model; il metodo cumulativo richiama anche il solutore lineare preesistente di GPCChecker.Concrete.")
table(["Grandezza","Convenzione"],[
["Coordinate","y = 0 all’interfaccia soletta acciaio; y positivo verso la soletta"],
["Azioni","N positivo a trazione; M positivo comprime le fibre superiori"],
["Piano di deformazione","ε(y) = ε₀ − κ y; ε dimensionless, κ in 1/mm nella libreria"],
["Trasporto del momento","M(yᵣ) = M₀ + N yᵣ; M₀ = −Σ σᵢ Aᵢ yᵢ"],
["Interfaccia","mm, MPa, kN, kNm; curve con κ in 1/m e ε in µε"],
["Deformazione della fibra","εmecc = εtot − εgetto − εimposta"],
["Azioni SLU","Gli incrementi N e M sono già combinati dall’utente; non si applicano ulteriori fattori alle azioni"]],[4.2,12.8])
h("Benchmark comune in cinque fasi")
p(f"Sezione ad H con soletta: soletta 3000 × 250 mm; piattabanda superiore 500 × 25 mm; anima 14 × 1800 mm; piattabanda inferiore 700 × 30 mm. Armature assenti per isolare la soluzione chiusa. Ea = {fmt(cum['Es'])} MPa, Ecm = {fmt(cum['Ec'],9)} MPa. Larghezze efficaci disattivate; N applicato a y = 0.")
table(["Fase","Stato","ΔN kN","ΔM kNm","φ","ψL","Δεcs µε"],[
["G1","Solo acciaio","−100","1500","0","1","0"],
["G2","Composta","−200","2000","2","1,1","0"],
["R1","Ritiro","0","0","1","0,55","−100"],
["Q","Composta","50","−500","0","1","0"],
["R2","Ritiro","0","0","2","0,55","−200"]],[1.1,3.2,2.4,2.5,1.2,1.2,5.4])
p("Il riferimento analitico integra aree, primi momenti e inerzie di elementi rettangolari e risolve un sistema 2 × 2. Il ritiro introduce εcs solo nel calcestruzzo. Le somme degli effetti sono calcolate direttamente nel test, senza richiamare i metodi di proprietà o tensione della libreria.")

# 3
page("Validazione del metodo cumulativo")
p("Per ogni situazione il metodo riparte dalla geometria lorda, calcola i contributi di tutte le fasi incluse, valuta le riduzioni di anima e piattabande dalle tensioni cumulate e ripete fino alla convergenza della sezione efficace. Tutti i contributi della situazione sono ricalcolati sulla stessa geometria efficace.")
p("Questo metodo non conserva lo stato deformativo costruttivo quando la sezione efficace cambia. Tale comportamento è una scelta del metodo, verificata dal test Class4_PriorSteelContributionIsRecomputedOnEachCommonEffectiveSection. Non deve essere interpretato come un integratore della storia del materiale.")
h("Riferimento geometrico del confronto")
p("Nelle fasi composte il solutore nativo integra le piattabande lungo la loro linea media: trascura l’inerzia locale b t³/12 delle piattabande, mantenendo il termine di trasporto. Il benchmark cumulativo usa lo stesso modello a parete sottile; quello storico lineare usa i rettangoli pieni. Un confronto iniziale con i rettangoli pieni mostrava, ad esempio, −1,47158640 MPa invece di −1,47163232 MPa al lembo superiore del CLS dopo G2. La differenza è modellistica e non viene nascosta aumentando la tolleranza.")
selected=[r for r in cum["Rows"] if r["Quantity"]=="Acciaio · intradosso"]
table(["Fase","Atteso MPa","Checker MPa","Scarto MPa"],[[names[i],fmt(r["Expected"],9),fmt(r["Actual"],9),fmt(abs(r["Actual"]-r["Expected"]),3)] for i,r in enumerate(selected)],[2,5,5,5])
p("Criterio del benchmark in cinque fasi: scarto assoluto ≤ 2 × 10⁻⁶ MPa per ciascun punto attivo di ciascuna fase. Le prove ordinarie comprendono carico nullo, trasporto del riferimento, due piastre inferiori, file di armatura opzionali, fessurazione per esclusione della soletta, φ/n, ritiro e iterazione di classe 4.")
picture(FIG/"fasi.png","Evoluzione di due tensioni significative nel benchmark cumulativo. Le ascisse sono stati successivi, non tempi fisici.",15.5)

# 4
page("Validazione dello storico lineare")
p("La fase attiva nuovi componenti senza tensioni nel piano deformato raggiunto al getto. Il riferimento di nascita è memorizzato per ogni fibra. Ogni incremento aggiunge le proprie azioni e deformazioni imposte; le tensioni pregresse restano memorizzate. Una fibra temporaneamente inattiva non porta risultanti ma conserva il proprio stato.")
p("Il rapporto di omogeneizzazione è n = Ea/Ecm · (1 + ψL φ). Nel metodo storico modifica il modulo dei nuovi incrementi di calcestruzzo. Non riscrive le tensioni già acquisite e non costituisce una legge viscosa dipendente dall’età. La forma del rapporto modulare è documentata nel materiale JRC di Davaine [3]; il trattamento incrementale dello storico è l’ipotesi specifica di questa implementazione.")
rows=linear["Rows"]
table(["Fase","ε₀ attesa µε","ε₀ Checker µε","κ attesa 1/m","κ Checker 1/m"],
      [[names[i],fmt(rows[2*i]["Expected"]*1e6,8),fmt(rows[2*i]["Actual"]*1e6,8),fmt(rows[2*i+1]["Expected"]*1000,8),fmt(rows[2*i+1]["Actual"]*1000,8)] for i in range(5)],[1.2,4,4,3.9,3.9])
p("Tolleranze: 2 × 10⁻¹² su ε₀, 2 × 10⁻¹⁴ 1/mm su κ e 2 × 10⁻⁷ MPa su ogni fibra attiva. Sono verificate tutte le fasi, non soltanto il risultato finale.")
h("Prove che distinguono una storia da una somma di carichi")
p("Un caso analitico con due componenti assiali di uguale EA applica 40 kN all’acciaio, attiva il secondo componente e applica altri 40 kN. Risultano 300 MPa nell’acciaio e 50 MPa nel secondo componente. Se entrambi sono attivi dall’inizio, a pari N finale, i valori sono 200 e 100 MPa. Lo scarico dopo il getto può quindi lasciare tensioni residue autoequilibrate anche con legami elastici.")
p("Sono inoltre provati getto a tensione nulla, ritiro libero, ritiro impedito internamente, ritiri multipli intercalati ai carichi, cambio del modulo senza nuove azioni, riattivazione, annullamento del calcolo e sequenze di oltre venti fasi. I riferimenti OpenSees elastici comprendono N–M, ritiro e getto dopo flessione.")
h("Classe 4")
p("Lo storico lineare può iterare le aree efficaci mantenendo il medesimo stato materiale consolidato durante tutti i tentativi. I test controllano convergenza ed equilibrio; il confronto esterno OpenSees descritto qui usa aree lorde fisse. Non è quindi una validazione esterna dell’intera evoluzione delle larghezze efficaci.")

# 5
page("Validazione dello storico non lineare")
p("La risposta è ottenuta integrando σᵢ Aᵢ e −σᵢ Aᵢ yᵢ. Due incognite, ε₀ e κ, soddisfano l’equilibrio N–M₀ mediante Newton, ricerca del passo e sottopassi. I tentativi iterativi non modificano la memoria consolidata; le variabili del materiale si accettano solo dopo l’equilibrio.")
table(["Materiale","Legge adottata","Memoria e limite"],[
["Carpenteria e armature","Bilineare con incrudimento isotropo","Deformazione plastica e accumulo; scarico elastico"],
["Calcestruzzo","Inviluppo tabulato Model","Risposta sull’inviluppo; nessun danno o isteresi ciclica"],
["Viscosità","Istantaneo φ = 0 oppure rifiuto di φ ≠ 0","Nessuna evoluzione viscosa temporale"]],[3.7,6.4,6.9])
h("Controllo analitico di carico e scarico")
p("Due fibre d’acciaio di area 100 mm² ciascuna, E = 200000 MPa, fy = 200 MPa, tangente post snervamento Et = 10000 MPa. N = 60 kN produce σ = 300 MPa ed ε = 0,011; scaricando N a zero rimane ε = 0,0095. I valori derivano direttamente dalla bilineare e dallo scarico con pendenza E.")
table(["Quantità","Atteso","Checker","Scarto","Tolleranza"],metric_rows(analytical[:2]),[4.4,3.1,3.1,3.2,3.2])
h("Confronto esterno della storia")
labels={"elastic_N_M":"N–M elastico","plastic_axial_cycle":"Ciclo assiale plastico","plastic_bending_cycle":"Ciclo di flessione plastica","plastic_N_M_cycle":"Ciclo N–M non proporzionale","composite_nonlinear":"Composta non lineare","free_composite_shrinkage":"Ritiro libero della composta","cast_then_shrink":"Getto dopo flessione e ritiro"}
table(["Storia","Stati","Max Δσ MPa","Max Δε"],[[labels[c["name"]],len(c["stages"]),fmt(history_errors[c["name"]][0],5),fmt(history_errors[c["name"]][1],5)] for c in history],[7.2,1.3,4.2,4.3])
p(f"Per ciascuno dei 29 stati si confrontano piano, tensione e deformazione di ogni fibra. Massimo scarto misurato: {fmt(max_hist,8)} MPa. Il confronto verifica il solutore e le leggi dichiarate, non le proprietà sperimentali dei materiali reali.")

# 6
page("Metodo momento curvatura")
p("A ogni punto viene imposta una curvatura κ. Il solutore ricerca la sola ε₀ necessaria a mantenere N costante; il momento è una risultante, non un carico assegnato. Lo schema corrisponde alla procedura di sezione a curvatura imposta descritta dagli esempi ufficiali OpenSees [4].")
p("L’avvio può essere vergine oppure riprendere una fase dello storico. Restano conservati εgetto, εimposta, stato plastico, componenti attivi e fattori elastici compatibili. Se viene richiesto un N diverso da quello iniziale, il precarico si applica a κ iniziale costante. Lo stato iniziale completo compare fra i punti della curva.")
code("εᵢ = ε₀ − κ yᵢ − εgetto,ᵢ − εimposta,ᵢ\nRₙ(ε₀) = Σ σᵢ(εᵢ, stato precedente) Aᵢ − Nassegnato = 0\nMᵣ = −Σ σᵢ Aᵢ yᵢ + N yᵣ")
h("Soluzioni chiuse di controllo")
p("Per il rettangolo 100 × 200 mm con E = 200000 MPa, I = b h³/12. Nel campo elastico M = EI κ. Per acciaio perfettamente plastico a N = 0, Mpl = fy b h²/4 e, dopo il primo snervamento, M = Mpl [1 − (κy/κ)²/3], con κy = fy/(E h/2). Il test usa 600 strisce e ammette uno scarto di 2 × 10⁻⁶ Mpl rispetto alla soluzione continua.")
table(["Quantità","Atteso","Checker","Scarto","Tolleranza"],metric_rows([analytical[2]]),[4.4,3.1,3.1,3.2,3.2])
picture(FIG/"momento_curvatura.png","Ciclo M–κ con acciaio bilineare e N = 0. I numeri indicano l’ordine dei sei stati; scarico e inversione mantengono la memoria.",16)
p("Il metodo può attraversare un plateau di momento in controllo di curvatura. Non garantisce il superamento di ogni biforcazione o punto limite: un ramo assiale non raggiungibile viene segnalato e non sostituito con uno stato non equilibrato.")

# 7
page("Metodo forza deformazione")
p("Si impongono ε alla quota yᵣ e una curvatura costante. Il piano è quindi noto: ε₀ = εᵣ + κ yᵣ. Le leggi costitutive restituiscono le tensioni, dalla cui integrazione si ottengono N e M. Non è richiesta l’inversione di una rigidezza assiale nulla: il metodo attraversa il plateau plastico N–ε.")
p("Curvatura costante non significa momento nullo. Nella sezione mista il momento risultante rappresenta la reazione necessaria a mantenere il vincolo imposto. Il grafico N–ε è accompagnato da M nella tabella e dai valori del piano, così il vincolo è esplicito.")
table(["Quantità","Atteso","Checker","Scarto","Tolleranza"],metric_rows(analytical[3:]),[4.4,3.1,3.1,3.2,3.2])
p("Il primo controllo usa il rettangolo elastico 100 × 200 mm: a ε = 0,001 si ha N = EA ε = 4 MN. Il secondo usa fy = 250 MPa e plasticità perfetta: a ε = 0,005 si ha N = A fy = 5 MN. Sono inoltre verificati scarico, inversione, κ non nulla e trasporto della quota di riferimento.")
picture(FIG/"forza_deformazione.png","Ciclo N–ε dell’acciaio con incrudimento isotropo. Le croci Checker coincidono con i punti OpenSees alla scala del grafico.",16)
h("Curva della singola fibra")
p("La vista σ–ε usa εmeccanica, non εtotale. La distinzione è essenziale dopo il getto o il ritiro: una fibra può avere una deformazione totale non nulla e tensione nulla. Il cursore dei due grafici seleziona lo stesso stato; i dati mostrano anche εgetto, εimposta ed εplastica.")

# 8
page("Confronto esterno delle nuove curve")
p("Il riferimento è generato da OpenSees 3.8.0 con FiberSection, zeroLengthSection e DisplacementControl. Gli script non importano Checker e non leggono suoi risultati. Stessa discretizzazione e stesse leggi sono dati comuni di ingresso; piani e stati sono risolti separatamente. Il confronto non misura l’errore comune di discretizzazione.")
curve_labels=["Mκ elastico eccentrico","Mκ ciclo plastico","Mκ con compressione","Mκ sezione composta","Nε ciclo plastico","Nε composta compressa"]
table(["Caso","Stati","Max Δσ MPa","Max ΔN N","Max ΔM Nmm"],[[curve_labels[i],len(c["points"]),fmt(actual[i]["maxStress"],5),fmt(actual[i]["maxForce"],5),fmt(actual[i]["maxMoment"],5)] for i,c in enumerate(curves)],[5.7,1.3,3.3,3.3,3.4])
p(f"Sono verificati 35 stati, incluse tensioni e deformazioni di ogni fibra. Massimo scarto tensionale: {fmt(max_curve,9)} MPa. Per i casi assiali la curvatura del riferimento esterno è zero; κ non nulla e yᵣ diverso da zero hanno confronti analitici separati.")
table(["Grandezza","Tolleranza assoluta","Quota relativa"],[
["ε₀ e deformazioni delle fibre","2 × 10⁻⁹","2 × 10⁻⁷ |atteso|"],
["κ","2 × 10⁻¹² 1/mm","2 × 10⁻⁷ |atteso|"],
["Tensioni","5 × 10⁻⁴ MPa","2 × 10⁻⁷ |atteso|"],
["N","0,02 N","2 × 10⁻⁷ |atteso|"],
["M","5 Nmm","2 × 10⁻⁷ |atteso|"]],[5.2,5.8,6])
h("Sensibilità dei percorsi plastici")
p("La campagna ha usato 200, 2000 e 20000 sottopassi per tratto. Il ciclo Mκ con inversione ha richiesto 200000 sottopassi per soddisfare le tolleranze prefissate senza allargarle. Il riferimento finale usa 20000 negli altri casi. Questi valori servono al benchmark, non sono impostazioni consigliate per ogni calcolo.")
p("Il numero di punti visualizzati e i sottopassi di integrazione sono distinti. Per un caso reale confrontare progressivamente punti, sottopassi e mesh; la convergenza del residuo di equilibrio da sola non prova l’indipendenza dal percorso discretizzato. Il confronto storico non lineare ha uno studio analogo, riportato nel README dei riferimenti.")

# Reproducible external model inputs
page("Dati dei modelli OpenSees")
p("I riferimenti OpenSees usano la sezione seguente, distinta dall’H alto 1855 mm del benchmark in cinque fasi. Il confronto si svolge sui medesimi punti di integrazione, con due punti di Gauss per striscia. Non usa la geometria calcolata da Checker come valore atteso.")
table(["Componente","b mm","h mm","y inferiore mm","Strisce"],[
["Piastra inferiore","200","20","−300","4"],["Anima","10","280","−280","28"],
["Piastra superiore","150","15","0","4"],["Soletta quando presente","600","100","15","32"]],[5.4,2.7,2.7,3.8,2.4])
p("Acciaio elastico: E = 200000 MPa. Acciaio non lineare: stesso E, fy = 250 MPa, Et = 2000 MPa e Hiso = E Et/(E − Et), Hkin = 0. Il CLS elastico usa E = 30000 MPa. Il CLS non lineare usa i punti della tabella seguente; il parametro InitialModulus = 30000 MPa serve al predittore, mentre la pendenza iniziale della tabella è 22500 MPa.")
table(["ε","−0,0035","−0,002","−0,001","0","0,01"],[["σ MPa","−30","−30","−22,5","0","0"]],[2.5,2.9,2.9,2.9,2.9,2.9])
h("Percorsi delle nuove curve")
table(["Caso","Vincolo","Target assoluti in ordine"],[
[curve_labels[0],"N = −100 kN","κ = 0,001; 0,005; 0; −0,005 1/m; yᵣ = −100 mm"],
[curve_labels[1],"N = 0","κ = 0,002; 0,01; 0,04; 0; −0,04; 0 1/m"],
[curve_labels[2],"N = −500 kN","κ = 0,002; 0,01; 0,03; 0; −0,03 1/m"],
[curve_labels[3],"N = −100 kN","κ = 0,001; 0,003; 0,008; 0,01; 0,003; 0 1/m"],
[curve_labels[4],"κ = 0","ε = 500; 2000; 6000; 4000; 0; −4000; 0 µε"],
[curve_labels[5],"κ = 0","ε = −250; −800; −1500; −2500; −3300; −1000; 0 µε"]],[5.2,3.4,8.4])
p("Per le curve Mκ il precarico assiale è applicato con rotazione bloccata; prima di liberare la rotazione il riferimento OpenSees sostituisce la reazione con il momento corrispondente. Durante la curva varia il momento mantenendo N. Per Nε la rotazione resta bloccata a zero. Il momento è riportato all’origine salvo il caso eccentrico, che usa yᵣ = −100 mm.")

page("Percorsi del confronto storico esterno")
p("Ogni riga elenca le risultanti totali dopo ciascuno stato, non gli incrementi. Il test Checker ricava gli incrementi per differenza. Il riferimento OpenSees mantiene lo stato dell’elemento tra i carichi, senza ricostruire il materiale plastico a ogni fase.")
table(["Storia","N totale kN","M₀ totale kNm"],[
[labels[c["name"]],"; ".join(fmt(s["n"]/1000) for s in c["stages"]),"; ".join(fmt(s["m"]/1e6) for s in c["stages"])] for c in history],[6.1,5.5,5.4])
p("Nel ritiro libero della composta entrambi i materiali sono elastici e il CLS riceve εcs = −200 µε. Nel caso con getto, prima si calcola il solo acciaio a M₀ = 80 kNm; il piano così ottenuto da OpenSees è il riferimento di nascita della soletta. Dopo il getto si applica il ritiro di −200 µε e poi la seconda coppia N–M della tabella.")
birth=history[-1]["birth"]
p(f"Piano di nascita esterno del caso con getto: ε₀ = {fmt(birth[0],10)}; κ = {fmt(birth[1],10)} 1/mm. La memoria è rappresentata nel riferimento elastico mediante InitStrainMaterial; la convenzione sottrae deformazione di nascita e deformazione imposta dalla deformazione totale.")
h("Criteri del confronto storico")
p("Su ε si usa una tolleranza assoluta di 2 × 10⁻¹⁰, su κ 2 × 10⁻¹² 1/mm e su σ 2 × 10⁻⁵ MPa, più 2 × 10⁻⁷ del valore atteso. Nei tre stati più sensibili dopo inversione plastica il pavimento di confronto, motivato dallo studio dei sottopassi, è 2 × 10⁻⁸ su ε e 0,004 MPa su σ. Gli scarti effettivi sono molto inferiori al limite nei riferimenti finali.")
p("Il riferimento storico usa 20000 sottopassi per fase e 200000 nel ciclo N–M non proporzionale. Il residuo OpenSees NormUnbalance è impostato a 0,001 nelle unità miste N e Nmm. Gli script si interrompono se OpenSees non converge e non salvano uno stato fallito come riferimento valido.")

# 9
page("Classe 4 e verifiche accessorie")
p("Il benchmark del pannello compresso usa il sottopannello 2 dell’esempio JRC EUR 22898 EN, capitolo 17, pagina stampata 200 [1]: b = 492 mm, t = 8 mm, fy = 235 MPa, ψ = 0,406. Il benchmark a taglio usa l’esempio di ponte composto JRC [2]: hw = 2720 mm, tw = 18 mm, a = 8000 mm, E = 210000 MPa e ν = 0,3. I valori pubblicati sono arrotondati.")
table(["Parametro","Pubblicato","Checker","Scarto","Tolleranza"],metric_rows(published),[4.4,3.1,3.1,3.2,3.2])
p("Questi confronti verificano coefficienti e riduzioni locali; non costituiscono un confronto esterno dell’intera sezione efficace di un ponte. Le curve non lineari restano lorde e non combinano la plasticità con la post instabilità delle piastre.")
table(["Controllo","Prove disponibili","Limite della validazione"],[
["Taglio dell’anima","Esempio JRC, rami di snellezza, pannelli quadrati e limiti","Verifica locale; nessuna torsione del cassone"],
["Pioli uniformi","Espressioni di resistenza, geometria ammessa, flusso V S/I e segni","Connessione completa; nessuno scorrimento"],
["Irrigidimenti e appoggi","Integrazione geometrica, eccentricità, soluzione secante, instabilità","Modello locale dichiarato; vincoli assegnati"],
["Saldature e soletta","Espressioni normative, minimi, campi di applicazione","Controlli analitici interni; niente confronto FEM dedicato"],
["N M V e fatica pioli","Soglie, inversioni, inviluppi dedicati, esclusione di dati mancanti","Accessori del cumulativo; non ricalcolati nei metodi storici"]],[3.8,6.9,6.3])
p("Le verifiche accessorie applicano i coefficienti del set normativo selezionato. Le curve usano legami caratteristici: il superamento o meno del limite della legge costitutiva non è un esito SLU di progetto.")

# 10
page("Interfaccia e conservazione dei risultati")
p("La nuova scheda Curve della sezione contiene i due tipi di curva, lo stato di partenza, il riferimento, il vincolo costante e l’incremento finale. La rampa parte dal valore iniziale e termina al valore iniziale più Δ; i segni possono essere positivi o negativi. La libreria accetta anche sequenze arbitrarie di valori assoluti, per cicli e inversioni.")
ui=ART/"ui-finale/curve_momento_curvatura.png"
if not ui.exists(): ui=ART/"ui-05/curve_momento_curvatura.png"
picture(ui,"Scheda delle curve in ANTHEA. Il cursore seleziona contemporaneamente stato di sezione e risposta della fibra.",17)
h("Controlli della presentazione")
p("Le prove WPF verificano la nuova scheda, Mκ e Nε, numero di punti, fibra selezionata, cursore sincronizzato, segno di N, vincolo di curvatura, ripresa dal getto, salvataggio delle preferenze, riapertura della scheda e segnalazione dei risultati da ricalcolare. È controllato anche che una curvatura piccola non venga visualizzata come zero.")
p("Il CSV delle curve contiene tutti i punti e tutte le fibre, inclusi stato di completamento, quota, tensione e deformazioni totale, al getto, imposta, meccanica e plastica. I punti non convergenti non vengono accettati. Se la curva si interrompe, rimane visibile l’ultimo stato valido con un messaggio esplicito.")
p("Il risultato di una curva è distinto da quello delle fasi. Cambiare geometria o materiale invalida l’esportazione del risultato precedente finché non viene ricalcolato. Le opzioni delle curve non cancellano quelle del metodo cumulativo o storico.")

# 11
page("Difetti aperti e approssimazioni")
p("La suite separata dei difetti noti è stata eseguita contro le DLL effettivamente distribuite da ANTHEA. I risultati seguenti non sono stati corretti in questa attività e non sono presentati come verifiche superate.")
table(["Caso","Esito","Effetto e perimetro"],[
["Costruttore H nullo","4 non superati","ReinforcedConcreteSection nello snapshot GPCModel genera un indice fuori intervallo. Il modulo H usa un profilo non nullo."],
["Inerzia del carico nullo","1 non superato","SolverInertia del cumulativo differisce fra azione nulla e non nulla: 69779283266,85039 contro 69776941198,49962 mm⁴ nel caso del test."],
["Riferimento ruotato","3 ignorati","Casi fuori dal percorso N–Mx con assi paralleli usato dall’adattatore"],
["Acciaio immerso nel CLS","4 ignorati","Caso nativo diverso dalla trave d’acciaio esterna alla soletta"],
["Asse neutro globale","3 ignorati","Diagnostica nativa esclusa dai campi storici a fibre"]],[4.3,3.1,9.6])
p("I dieci casi ignorati non forniscono una nuova misura del difetto: restano non verificati in questa esecuzione. Il dettaglio dei nomi e dei motivi è conservato nel TRX difetti-noti e nei test sorgente.")
h("Approssimazioni esplicite")
p("Le due piastre inferiori sono sostituite nel calcolo H da un rettangolo con la stessa area e lo stesso spessore totale. Se hanno larghezze diverse non conserva esattamente baricentro e inerzia delle piastre reali; i test ne quantificano l’effetto. Nei metodi storici i rettangoli e le barre sono integrati con due punti per striscia o componente equivalente.")
p("Le tensioni disegnate alle facce nel metodo non lineare usano le fibre vicine, senza estrapolazioni oltre il legame costitutivo. I massimi riportati sono massimi sui punti di integrazione. Occorre raffinare la mesh se gli estremi o la posizione dell’asse neutro sono sensibili.")
p("Restano fuori dal campo validato: viscosità dipendente dall’età, connessione parziale, danno ciclico del CLS, instabilità globale, torsione di cassoni, interazione post critica plastica di classe 4 e risposta tridimensionale. L’API a fibre è generica, ma questa campagna non valida una famiglia completa di cassoni.")

# 12
page("Inventario delle prove numeriche")
p(f"Suite CompositeBridge: {composite['passed']} casi superati. Suite BridgeAudit ordinaria: {audit['passed']} casi superati. Totale {total}. I casi parametrizzati sono contati singolarmente; i sei test delle curve OpenSees comprendono complessivamente 35 stati. Il confronto storico OpenSees conta 29 test, uno per stato.")
table(["Classe di test CompositeBridge","Casi"],sorted(cc.items()),[14,3])
p("Il dossier aggiunge quattro prove con emissione dei valori numerici: cinque fasi cumulative, cinque fasi storiche lineari, benchmark pubblicati e soluzioni chiuse non lineari / curve. La loro soluzione attesa è analitica e scritta nello sviluppo; non è etichettata come validazione di terzi.")
h("Criteri di accettazione")
p("Ogni prova confronta grandezze significative e verifica invarianti fisiche: segni, equilibrio, conservazione dello stato, compatibilità del getto, ordine delle fasi e comportamento in scarico. I confronti con la stessa libreria o con snapshot precedenti sono regressioni; verificano la stabilità del comportamento, non l’esattezza fisica di per sé.")

# 13
page("Inventario delle prove di integrazione")
table(["Classe di test BridgeAudit ordinaria","Casi"],sorted(ac.items()),[14,3])
p("Gli otto casi BridgeResponseIntegrationTests verificano unità, rampa, origine storica, CSV, valori fuori campo e conservazione delle preferenze. Sono stati rieseguiti anche contro la DLL finale copiata in ANTHEA.")
h("Regressioni della vista")
p("La suite WPF del modulo controlla geometria, tensioni, proprietà, ritiro, taglio, pioli, dettagli, storico e vista separata. La nuova scheda dispone di una prova dedicata e della stessa prova richiamata dal collaudo completo del modulo. Le schermate e i conteggi dei controlli sono archiviati negli artefatti di questa revisione.")
p("Nel banco di prova automatico sono soppressi i tooltip: la creazione e chiusura rapida delle finestre sotto il puntatore causava un riferimento WPF a un handle già distrutto. La soppressione è limitata alla modalità smoke; i tooltip dell’applicazione restano attivi.")

# 14
page("Riproducibilità e identificazione")
p("I test numerici risiedono nel repository Checker. Documentazione, script di impaginazione, schermate, log e dati del dossier risiedono sotto supporto in ANTHEA. OpenSees e Matplotlib sono strumenti di prova e documentazione; non sono dipendenze dell’applicazione distribuita.")
table(["Elemento","Posizione"],[
["Calcolo delle fasi","Checker/GPCChecker.CompositeBridge/History"],
["Curve della sezione","BridgeSectionResponseAnalysis e HBridgeSectionResponse"],
["Test e riferimenti esterni","Checker/GPCChecker.Test.CompositeBridge/Validation"],
["Audit applicativo","Checker/GPCChecker.Test.BridgeAudit"],
["Evidenze di questa revisione","ANTHEA/supporto/artefatti/ponte_curve_validazione"],
["Generatore del dossier","ANTHEA/supporto/scripts/Build-BridgeValidation.py"]],[5.2,11.8])
h("Ripetizione delle prove")
code("dotnet test GPCChecker.Test.CompositeBridge -c Release\n  -p:BridgeLibraryDir=<ANTHEA>/lib/Checker\n\ndotnet test GPCChecker.Test.BridgeAudit -c Release\n  --filter \"TestCategory!=KnownBug&TestCategory!=ConstructorRegression\"\n\ndotnet test GPCChecker.Test.BridgeAudit -c Release\n  --filter \"TestCategory=KnownBug|TestCategory=ConstructorRegression\"")
p("Impostare BRIDGE_VALIDATION_OUTPUT a una cartella di artefatti per emettere i JSON del dossier. I riferimenti sono rigenerabili con Python 3.12 e openseespy 3.8.0.0: opensees_reference.py per lo storico; opensees_curves.py con 20000 sottopassi base per le curve. Nel ciclo Mκ più sensibile lo script moltiplica il valore base per dieci.")
h("Impronte SHA256")
for label,path in [
("GPCChecker.CompositeBridge.dll",ROOT/"lib/Checker/GPCChecker.CompositeBridge.dll"),
("GPCChecker.Concrete.dll",ROOT/"lib/Checker/GPCChecker.Concrete.dll"),
("GPCModel.dll",ROOT/"lib/Checker/GPCModel.dll"),
("Riferimento OpenSees storico",VAL/"opensees-3.8.0.json"),
("Riferimento OpenSees curve",VAL/"opensees-curves-3.8.0.json")]:
    p(label);code(sha(path))

# 15
page("Riferimenti e valutazione finale")
p("Le edizioni richiamate sono quelle implementate dal modulo e dai benchmark. Questo dossier verifica l’implementazione dichiarata; non sostituisce il controllo delle norme e degli allegati nazionali applicabili alla singola commessa.")
refs=[
("[1] JRC EUR 22898 EN, 2007","Commentary and worked examples to EN 1993 1 5. Capitolo 17, sottopannello 2, pagina stampata 200, pagina PDF 214.","https://eurocodes.jrc.ec.europa.eu/sites/default/files/2021-12/EUR22898EN.pdf"),
("[2] JRC, 2012","Bridge Design to Eurocodes Worked examples. Capitolo 6, verifica a taglio, pagina stampata 130.","https://eurocodes.jrc.ec.europa.eu/sites/default/files/2022-06/Bridge_Design-Eurocodes-Worked_examples-main_only.pdf"),
("[3] L. Davaine, JRC, 2013","Worked examples on bridge design with Eurocodes. Rapporti modulari, slide 11 e 14.","https://eurocodes.jrc.ec.europa.eu/sites/default/files/2022-06/D1.8Davaine.pdf"),
("[4] OpenSees Examples Manual","Example 2, Moment Curvature Analysis of a RC Section. Procedura con controllo della rotazione della sezione.","https://opensees.berkeley.edu/OpenSees/manuals/ExamplesManual/HTML/3886.htm"),
("[5] OpenSeesPy","Documentazione del materiale Hardening e dei parametri di incrudimento.","https://openseespydoc.readthedocs.io/en/latest/src/Hardening.html"),
("[6] OpenSeesPy","Documentazione del materiale ElasticMultiLinear.","https://openseespydoc.readthedocs.io/en/stable/src/ElasticMultiLinear.html"),
("[7] Ministero delle Infrastrutture e dei Trasporti","D M 17 gennaio 2018. Aggiornamento delle Norme tecniche per le costruzioni. GU 42 del 20 febbraio 2018, supplemento ordinario 8.","https://www.gazzettaufficiale.it/eli/id/2018/2/20/18A00716/sg")]
for title,desc,url in refs:
    para=p(title+" · "+desc); para.paragraph_format.keep_with_next=True
    para=p(url);para.paragraph_format.space_after=Pt(9)
    for run in para.runs:run.font.size=Pt(8.5)
h("Giudizio sul campo provato")
p("I metodi risultano numericamente coerenti con i riferimenti dichiarati nei casi provati. L’evidenza più forte è il confronto su piani e fibre con un solutore esterno; quella più limitata riguarda le parti verificate solo con formule locali o regressioni. Prima di estendere l’uso occorrono confronti mirati per nuove geometrie, legami, percorsi, instabilità e vincoli. I difetti aperti e i limiti di modello devono restare associati ai risultati.")

target=OUT/"ANTHEA_Validazione_Sezione_Ponte_Rev01.docx"
doc.save(target)
summary={"ordinary_passed":total,"composite_passed":int(composite["passed"]),"audit_passed":int(audit["passed"]),
         "known_failed":int(bugs["failed"]),"known_skipped":known_skipped,"history_states":29,"curve_states":35,
         "history_max_stress_MPa":max_hist,"curve_max_stress_MPa":max_curve,"document":str(target),
         "composite_dll_sha256":sha(ROOT/"lib/Checker/GPCChecker.CompositeBridge.dll")}
(ART/"esiti.json").write_text(json.dumps(summary,indent=2),encoding="utf-8")
print(json.dumps(summary,indent=2))
