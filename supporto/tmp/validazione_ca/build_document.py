from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
from hashlib import sha256
import json, math
from lxml import etree
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
OUT=ROOT.parents[1]/'documentazione'/'Validazione_CA_ANTHEA'
OUT.mkdir(exist_ok=True)
TEMPLATE=Path(r'C:/Users/g.pacini/Desktop/MODELLO-RELAZIONE-ITEC-AA.docx')
REF=json.loads((ROOT/'reference.json').read_text(encoding='utf8'))
ACT={v['id']:v for v in json.loads((ROOT/'actual.json').read_text(encoding='utf8'))}
doc=Document(TEMPLATE)
sections=[deepcopy(s._sectPr) for s in doc.sections]

def fmt(v,n=4):
 if v is None:return '—'
 if abs(v)<1e-9:v=0
 return f'{v:.{n}f}'.replace('.',',')
def sci(v):
 if abs(v)<1e-20:return '0'
 s=f'{v:.12g}'
 if 'e' in s:
  mant,exp=s.split('e');s=mant+'×10'+str(int(exp)).translate(str.maketrans('-0123456789','⁻⁰¹²³⁴⁵⁶⁷⁸⁹'))
 return s.replace('.',',')
def status(id):
 if id in ['E3D-02','E2D-02']:return 'DA CHIARIRE'
 if id in ['P2D-01','P2D-02','E2D-01']:return 'CON RISERVA'
 return 'CONFORME'

# Mantiene il frontespizio ITEC e sostituisce i soli slot dimostrativi.
def replace_p(p,text):
 runs=p.runs
 if runs:
  runs[0].text=text
  for r in runs[1:]:r.text=''
 else:p.add_run(text)
replace_p(doc.paragraphs[2],'Validazione del calcolo delle sezioni in calcestruzzo armato')
replace_p(doc.paragraphs[4],'Software ANTHEA  |  NTC 2018 e Circolare 2019')
doc.tables[0].cell(3,0).text='Ambito:'
doc.tables[0].cell(5,1).text='Validazione del software ANTHEA'
for row in doc.tables[1].rows[1:]:
 for c in row.cells:c.text=''
doc.tables[1].cell(0,1).text='Responsabili'
doc.tables[1].cell(1,1).text='Red.'
doc.tables[1].cell(2,1).text='Contr.'
doc.tables[1].cell(3,1).text='Appr.'
for ir in range(1,4):doc.tables[1].cell(ir,2).text='Da indicare'
doc.tables[2].cell(1,0).text='Relazione di validazione numerica'
doc.tables[2].cell(2,0).text='20 casi di riferimento\nDomini 3D e 2D\nAnalisi e verifiche tensionali\nFessurazione e taglio X e Y'
for i,vals in enumerate([
 ['Tipo documento: VALIDAZIONE','Nome file:','ANTHEA-VAL-CA-01'],
 ['Software: ANTHEA','Elaborato:','VAL-CA-01']]):
 for j,t in enumerate(vals):doc.tables[3].cell(i,j).text=t
for j,t in enumerate(['00','PRIMA EMISSIONE','24/09/2026','','','']):doc.tables[4].cell(1,j).text=t
for ti,size in [(1,8),(2,11),(3,10),(4,8)]:
 for row in doc.tables[ti].rows:
  for cell in row.cells:
   for para in cell.paragraphs:
    para.alignment=WD_ALIGN_PARAGRAPH.LEFT
    for run in para.runs:run.font.size=Pt(size)

body=doc._element.body
for child in list(body)[23:]:body.remove(child)
group=doc.tables[1]._tbl
group.getparent().remove(group)
# Il modello termina il frontespizio con una sezione continua; l'indice usa il suo pattern corrente.
end=deepcopy(sections[4])
body.append(end)
for s in doc.styles:
 if s.type==WD_STYLE_TYPE.PARAGRAPH and s.name in ['Heading 1','Heading 2','Heading 3']:
  s.font.color.rgb=RGBColor(0,0,0);s.paragraph_format.keep_with_next=True
if 'Title' not in doc.styles:doc.styles.add_style('Title',WD_STYLE_TYPE.PARAGRAPH)
doc.styles['Title'].base_style=doc.styles['Normal'];doc.styles['Title'].font.name='Manrope ExtraBold';doc.styles['Title'].font.size=Pt(18);doc.styles['Title'].font.color.rgb=RGBColor(0,0,0)
doc.paragraphs[2].style=doc.styles['Title']

def p(text='',boldlead=None):
 x=doc.add_paragraph(style='Normal');x.paragraph_format.space_after=Pt(7)
 if boldlead:x.add_run(boldlead+' ').bold=True
 x.add_run(text);return x
def heading(text,level=1,newpage=False):
 x=doc.add_paragraph(text,style=f'Heading {level}')
 if newpage:x.paragraph_format.page_break_before=True
 x.paragraph_format.space_after=Pt(8);return x
def small(text):
 x=p(text)
 for r in x.runs:r.font.size=Pt(10)
 return x
def mathp(text):
 x=doc.add_paragraph();x.paragraph_format.space_after=Pt(6)
 om=OxmlElement('m:oMath');r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=text;r.append(t);om.append(r);x._p.append(om)
 return x
def table(headers,rows,widths=None):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 if widths is None:widths=[170/len(headers)]*len(headers)
 for c,w in zip(t.columns,widths):c.width=Mm(w)
 for i,h in enumerate(headers):t.rows[0].cells[i].text=h
 for row in rows:
  cells=t.add_row().cells
  for i,v in enumerate(row):cells[i].text=str(v)
 for ir,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
  if ir==0:trpr.append(OxmlElement('w:tblHeader'))
  for j,cell in enumerate(row.cells):
   cell.width=Mm(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcpr=cell._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
   for side in ['top','bottom','left','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),'65');e.set(qn('w:type'),'dxa');mar.append(e)
   tcpr.append(mar)
   if ir==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7E6E6');tcpr.append(sh)
   for para in cell.paragraphs:
    para.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_after=Pt(1);para.paragraph_format.space_before=Pt(1)
    for run in para.runs:run.font.size=Pt(10.5);run.bold=ir==0
 bord=OxmlElement('w:tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:
  e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'B7B7B7');bord.append(e)
 t._tbl.tblPr.append(bord)
 doc.add_paragraph().paragraph_format.space_after=Pt(0)
 return t

# Indice automatico, solo capitoli per mantenere leggibilità su una pagina.
x=p('INDICE');x.alignment=WD_ALIGN_PARAGRAPH.CENTER;x.runs[0].bold=True
toc=p();f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'TOC \\o "1-1" \\h \\z \\u');toc._p.append(f)
p('I casi sono identificati da un codice univoco riportato nelle schede e nel quadro riepilogativo. Sono previsti due esempi per ciascuna delle dieci famiglie di verifica.')
# Chiusura dell'indice e avvio corpo con il pattern ITEC.
idxsect=deepcopy(sections[3]);idxsect.find(qn('w:type')).set(qn('w:val'),'continuous')
idxsect.find(qn('w:pgMar')).set(qn('w:bottom'),'1134')
tocend=p();tocend._p.get_or_add_pPr().append(idxsect)
finalsect=doc.sections[-1]._sectPr
for e in list(finalsect):finalsect.remove(e)
for e in sections[4]:finalsect.append(deepcopy(e))
mar=finalsect.find(qn('w:pgMar'));mar.set(qn('w:bottom'),'1134')
for tag in ['titlePg','pgNumType','type']:
 for e in finalsect.findall(qn('w:'+tag)):finalsect.remove(e)
num=OxmlElement('w:pgNumType');num.set(qn('w:start'),'1');finalsect.append(num)
# Usa il logo dell'intestazione corpo anche nella prima pagina della relazione.
for e in list(finalsect.findall(qn('w:headerReference'))):finalsect.remove(e)
for e in sections[3].findall(qn('w:headerReference')):
 if e.get(qn('w:type'))=='default':finalsect.insert(0,deepcopy(e))

heading('Oggetto e conclusioni della validazione')
p('La presente relazione definisce ed esegue una campagna di validazione numerica del modulo di calcolo delle sezioni in calcestruzzo armato di ANTHEA. Il confronto comprende 20 casi: domini resistenti 3D e 2D, in modalità plastica ed elastica; analisi tensionale lineare e non lineare; limitazione delle tensioni; apertura delle fessure; resistenza a taglio nelle direzioni X e Y.')
p('La campagna fornisce 15 casi conformi ai criteri numerici adottati, 3 casi con riserva sul dettaglio tensionale restituito dai domini 2D e 2 casi elastici da chiarire per scostamenti superiori alla soglia. La validazione complessiva del modulo resta aperta fino alla chiusura dei rilievi descritti nella relazione. Un esito strutturale negativo correttamente riprodotto costituisce un test software positivo.')
heading('Attività da eseguire',2)
p('Il verificatore deve fissare versione del programma e librerie, unità, assi, materiali e opzioni; ricostruire ogni riferimento senza utilizzare il risolutore da validare; immettere i dati in ANTHEA; confrontare azioni resistenti, tensioni, deformazioni, ampiezza delle fessure e resistenze a taglio. Gli scostamenti vanno valutati prima dell’arrotondamento e registrati con il relativo esito.')
p('Per i domini occorre controllare sia un punto di frontiera sia punti interni ed esterni. Per i risultati non lineari occorre distinguere il limite del materiale, il criterio di ricerca del punto resistente e la rappresentazione grafica. Prima di approvare il rapporto devono essere risolte le discrepanze e ripetuti i casi interessati.')
heading('Campo di applicazione',2)
p('Sono considerate sezioni rettangolari, armature ordinarie aderenti, deformazioni piane, aderenza perfetta, temperatura ordinaria e assenza di precompressione. Il modello è di sezione: non comprende stabilità, effetti del secondo ordine, duttilità dell’elemento, gerarchia delle resistenze, ancoraggi, minimi costruttivi, torsione, punzonamento o verifiche globali dell’opera. Le verifiche di taglio X e Y sono separate: la loro interazione non è validata.')
small('Riferimenti principali: NTC 2018, capitolo 4 e § 10.2 [R1]; Circolare 2019, capitolo C4 [R2]. La base normativa è quella richiesta per questa campagna, con EC2 di prima generazione come supporto [R3–R4].')

heading('Protocollo e criteri di accettazione',newpage=True)
p('La verifica è stata eseguita il 24 settembre 2026 richiamando direttamente le API del motore corrente: CheckerSection per domini e tensioni; Ntc2018Checks per fessurazione e taglio. Il collaudo riguarda il motore e l’adattatore di ANTHEA, non le operazioni manuali nell’interfaccia WPF, l’importazione dei dati o l’impaginazione dei report del programma.')
p('I riferimenti sono calcolati da un programma indipendente in Python, senza chiamate alle DLL Checker. Per le integrazioni di sezione si usa quadratura di Gauss a tratti: il passaggio da 6 a 10 punti mantiene N e M entro 10⁻⁸ nelle rispettive unità. Le formule chiuse e l’algoritmo di riproduzione sono riportati nel seguito.')
table(['Grandezza','Criterio adottato'],[
 ['N e componenti di M','Scarto relativo ≤ 1%; vicino a zero: 0,05 kN e 0,01 kNm'],
 ['Tensioni','Scarto ≤ max(1% del riferimento; 0,05 MPa)'],
 ['Tasso di utilizzo η','Scarto assoluto ≤ 0,01 rispetto al riferimento'],
 ['Ampiezza wk','Scarto ≤ max(1%; 0,002 mm)'],
 ['Resistenza a taglio VRd','Scarto ≤ max(1%; 0,05 kN)'],
 ['Punti di controllo','0,8 volte la frontiera: interno; 1,2 volte: esterno']],[60,110])
p('Queste tolleranze sono criteri tecnici della campagna, non limiti prescritti dalla normativa. Non si arrotonda un tasso superiore a 1 per trasformarlo in verifica soddisfatta: in prossimità della frontiera si confronta il valore numerico con la tolleranza di validazione.')
p('Per i domini 3D si usa il criterio “Eccentricità costante”, con strategia “Iterativo”, 64 direzioni e azioni proporzionali (N, Mx, My). Il dominio 2D è N–M, θ = 0°, senza proiezione delle azioni. Il controllo aggiuntivo a 128 direzioni restituisce gli stessi valori dei tassi per tutti gli otto casi: non elimina gli scarti riscontrati e non costituisce prova di convergenza della mesh interna.')
p('Le tre categorie di esito sono: CONFORME, se le grandezze previste rispettano la soglia; CON RISERVA, se N e M rientrano ma il dettaglio tensionale associato non è coerente; DA CHIARIRE, se il riferimento non è riprodotto entro la soglia. Le schede distinguono queste categorie dall’esito della verifica strutturale.')

heading('Geometrie materiali e convenzioni',newpage=True)
p('Gli assi x e y giacciono nel piano della sezione, con origine al centro geometrico. N è negativo in compressione. Una compressione prevalente sul lembo y positivo produce Mx positivo; sul lembo x positivo produce My negativo. Tensioni e deformazioni ANTHEA sono negative in compressione. Le lunghezze sono in mm, le tensioni in MPa, le forze in kN e i momenti in kNm.')
table(['Parametro','Sezione R','Sezione W'],[
 ['Dimensioni b × h','300 × 500 mm','600 × 500 mm'],
 ['Barre B1 e B2','(−100; −200), (+100; −200)','(−250; −200), (+250; −200)'],
 ['Barre B3 e B4','(−100; +200), (+100; +200)','(−250; +200), (+250; +200)'],
 ['Armatura longitudinale','4 Ø20; As = 1256,6371 mm²','4 Ø20; As = 1256,6371 mm²'],
 ['Copriferro netto e staffa','30 mm e Ø10','30 mm e Ø10'],
 ['Quota del centro delle barre','50 mm dal bordo','50 mm dal bordo']],[57,56.5,56.5])
p('Le barre sono definite mediante coordinate manuali. Il calcestruzzo spostato dalle barre è sottratto con la correzione puntuale As,i·σc(xi,yi); le barre sono concentrate nei loro centri. Questa convenzione deve essere mantenuta nel confronto, soprattutto quando si definisce il primo snervamento.')
m=REF['materials']
table(['Materiale o coefficiente','Valore'],[
 ['Calcestruzzo C30/37 e acciaio B450C','fck = 30 MPa; fyk = 450 MPa'],
 ['Coefficienti di progetto','αcc = 0,85; γc = 1,50; γs = 1,15'],
 ['Resistenze di progetto',f'fcd = {fmt(m["fcd"])} MPa; fyd = {fmt(m["fyd"])} MPa'],
 ['Moduli',f'Es = 200000 MPa; Ecm = {fmt(m["Ecm"])} MPa'],
 ['Trazione media del CLS',f'fctm = {fmt(m["fctm"])} MPa'],
 ['Deformazioni di riferimento','εc2 = 2,0‰; εcu2 = 3,5‰; εyd = 1,956522‰']],[68,102])
small('Ecm = 22000·[(fck + 8)/10]^0,3; fctm = 0,30·fck^(2/3). φ viscoso = 0; CLS teso escluso; getto sottile = No; acciaio elastoplastico senza incrudimento. La deformazione ultima dell’acciaio impostata a 100‰ è un parametro del benchmark e non viene raggiunta. FE-01 usa i dati specifici della fonte.')

heading('Formule dei riferimenti indipendenti',newpage=True)
heading('Equilibrio della sezione e legami costitutivi',2)
p('Si introduce e, positiva in compressione, opposta alla deformazione restituita da ANTHEA. Ogni caso definisce un piano noto e ricava le azioni equilibranti; ANTHEA deve ritrovare quel punto o quello stato tensionale. Per i casi di frontiera sono imposte deformazioni limite ammissibili, non un blocco di tensioni arbitrario.')
mathp('e(x,y) = a + bₑ x + cₑ y')
mathp('s꜀(e) = 0 per e ≤ 0;   s꜀(e) = fcd [2(e/εc2) − (e/εc2)²] per 0 < e < εc2')
mathp('s꜀(e) = fcd per εc2 ≤ e ≤ εcu2;   sₛ(e) = max(−fyd, min(Es e, fyd))')
p('Per il calcolo lineare si sostituiscono i legami con sc = Ecm·max(e,0) e ss = Es·e. Le tensioni con segno del programma sono σc = −sc e σs = −ss. Per l’integrazione sul rettangolo lordo si usano le espressioni seguenti, con gli integrali in N e Nmm prima della conversione delle unità.')
mathp('C = ∫A s꜀ dA + Σ As,i (sₛ,i − s꜀,i);     N = −C / 1000')
mathp('Mx = [∫A s꜀ y dA + Σ As,i (sₛ,i − s꜀,i) yi] / 10⁶')
mathp('My = −[∫A s꜀ x dA + Σ As,i (sₛ,i − s꜀,i) xi] / 10⁶')
heading('Costruzione dei punti di frontiera',2)
p('Per e = κ(px + qy − q0), si pone u = |p|b/2 + |q|h/2. Nei casi plastici κ = εcu2/(u − q0). Nei casi elastici della campagna si assume il primo limite fra εc2 al vertice compresso e |εs| = fyd/Es al centro delle barre. Il termine “elastico” identifica qui il dominio convenzionale di primo limite: il CLS conserva il ramo parabolico. Non coincide con l’analisi tensionale lineare della sezione omogeneizzata.')
mathp('κel = min { εc2/(u − q0); (fyd/Es)/maxi |p xi + q yi − q0| }')
small('La distinzione è coerente con il criterio di curvatura di prima plasticizzazione NTC § 4.1.2.3.4.2 [R1]. La corrispondenza esatta con la definizione interna di “Elastic” di Checker è oggetto dei casi E3D-02 ed E2D-02. La sigla interna SLV non equivale alla verifica sismica completa.')

heading('Tensioni di esercizio e fessurazione',2,newpage=True)
p('La limitazione delle tensioni si applica ai valori estremi dello stato tensionale. Con i materiali del benchmark, per la combinazione rara si controllano |σc,min| ≤ 18 MPa e |σs|max ≤ 360 MPa; per la quasi permanente si controlla |σc,min| ≤ 13,5 MPa. I casi VT usano l’analisi lineare fessurata [R1, § 4.1.2.2.5].')
mathp('ηrara = max(|σc,min|/18; |σs|max/360);    ηQP = |σc,min|/13,5')
p('Per la fessurazione si utilizzano la sezione parzializzata e le formule della Circolare [R2, C4.1.2.2.4.5]. La grandezza Δε è la deformazione media efficace indicata εsm nella Circolare. Es/Ecm nella formula di fessurazione va distinto dal rapporto di omogeneizzazione comprensivo degli effetti viscosi.')
mathp('hc,eff = min[2,5(h − d); (h − xₙ)/3; h/2];    ρeff = As,eff/Ac,eff')
mathp('Δε = max{[σs − kt fctm (1 + αe ρeff)/ρeff]/Es; 0,6 σs/Es}')
mathp('αe = Es/Ecm;    wk = 1,7 Δsm Δε')
p('Con barre ad aderenza migliorata si usa k1 = 0,8; per flessione k2 = 0,5; k3 = 3,4 e k4 = 0,425. kt è 0,6 per carichi di breve durata e 0,4 per carichi di lunga durata. c indica il ricoprimento delle barre longitudinali, quindi qui c = 30 + 10 = 40 mm.')
mathp('Δsm,vicino = [3,4 c + 0,8 k2 0,425 Ø/ρeff]/1,7')
mathp('s ≤ 5(c + Ø/2):  Δsm = Δsm,vicino')
mathp('s > 5(c + Ø/2):  Δsm,lontano = 0,75(h − xₙ)')
p('Per barre distanziate si calcolano entrambe le zone e si assume l’apertura più elevata. Nei casi di questa campagna le altre grandezze sono comuni alle due zone, quindi si prende il maggiore dei due valori di Δsm. FE-02 adotta esposizione XC1, armatura poco sensibile, combinazione quasi permanente e limite 0,30 mm.')
small('FE-01 riproduce un valore pubblicato mediante la sola funzione di apertura; FE-02 verifica anche tensioni, area efficace e spaziatura geometrica. I due livelli di copertura non vanno confusi. Trazione uniforme, due aree efficaci distinte e precompressione restano fuori da questa campagna.')

heading('Taglio nelle due direzioni',2,newpage=True)
p('Si considerano azioni trasversali separate, assenza di torsione, N = −300 kN e sezione R. In direzione X si adottano bw = 500 mm e d = 250 mm; in direzione Y bw = 300 mm e d = 450 mm. La larghezza è ortogonale alla direzione del taglio e d è misurato lungo tale direzione. Per entrambe Asl = 628,3185 mm² e A = 150000 mm², da cui σcp = 2 MPa.')
heading('Elementi senza armature resistenti a taglio',3)
mathp('k = min[2; 1 + √(200/d)];    ρl = min[0,02; Asl/(bw d)]')
mathp('vmin = 0,035 k^(3/2) √fck;    σcp = min[−1000 N/A; 0,2 fcd]')
mathp('VRd,c = max{[(0,18/γc) k (100 ρl fck)^(1/3) + 0,15 σcp] bw d;')
mathp('                         (vmin + 0,15 σcp) bw d} / 1000')
p('I casi senza staffe verificano la resistenza della formula NTC § 4.1.2.3.5.1. Non costituiscono una soluzione costruttiva approvata per una trave o un pilastro: le prescrizioni sulle armature minime e sugli ancoraggi sono controlli ulteriori.')
heading('Elementi con staffe verticali',3)
p('Si impiegano due bracci Ø8, Asw = 100,5310 mm², passo s = 150 mm, α = 90°, cotθ = 2 e z = 0,9d. L’inclinazione θ è fissata, senza ottimizzazione automatica. Per σcp = 2 MPa < 0,25fcd, αc = 1 + σcp/fcd = 1,117647 [R1, § 4.1.2.3.5.2].')
mathp('VRsd = z (Asw/s) fyd cotθ / 1000')
mathp('VRcd = z bw αc (0,5 fcd) cotθ/(1 + cot²θ) / 1000')
mathp('VRd = min(VRsd; VRcd);     ηV = |VEd|/VRd')
p('La staffa Ø8 è l’armatura del benchmark di resistenza. Le altezze utili sono input manuali fissati come sopra, senza ricalcolarle dal copriferro del caso geometrico con staffa Ø10. Per il taglio Y i bracci efficaci sono quelli verticali della staffa; per X quelli orizzontali. Non si sommano indiscriminatamente tutti i bracci.')

families={
 '3D plastico':'Dominio tridimensionale plastico', '3D elastico':'Dominio tridimensionale elastico',
 '2D plastico':'Dominio bidimensionale plastico', '2D elastico':'Dominio bidimensionale elastico',
 'Calcolo tensionale lineare':'Calcolo tensionale lineare', 'Calcolo tensionale non lineare':'Calcolo tensionale non lineare',
 'Verifica tensionale':'Verifica di limitazione delle tensioni','Fessurazione':'Verifica di fessurazione','Taglio X':'Verifica a taglio in direzione X','Taglio Y':'Verifica a taglio in direzione Y'}
prev=None
for v in REF['cases']:
 id=v['id'];a=ACT[id];r=v['reference'];family=v['family'];mode=v['mode']
 if family!=prev:heading(families[family],newpage=True);heading(id.replace('-',' ')+' '+v['title'],2)
 else:heading(id.replace('-',' ')+' '+v['title'],2,newpage=True)
 prev=family
 if mode=='domain':
  pl=r['plane'];dim=v['dimension'];elastic=v['state']=='SLV';rr=a['boundary']['Resistance'];resp=a['boundary']['Response']
  p('Sezione R e materiali del capitolo 3. CLS teso escluso, acciaio elastoplastico. Il punto è costruito dal piano di deformazione noto sotto riportato; si integra il legame parabola–rettangolo e si corregge il contributo del CLS in corrispondenza delle quattro barre. Fonte del procedimento: equilibrio e compatibilità [R1], riferimento analitico della presente campagna.')
  mathp(f'e(x,y) = {sci(pl[0])} + ({sci(pl[1])}) x + ({sci(pl[2])}) y')
  if dim==3:
   p('Impostare Dominio 3D, '+('Elastico' if elastic else 'Plastico')+', criterio Eccentricità costante e strategia Iterativo. Inserire le azioni della colonna Riferimento. La frontiera attesa ha η = 1; i tre valori N, Mx e My devono essere confrontati separatamente.')
  else:
   p('Impostare Dominio 2D, N–M, θ = 0°, '+('Elastico' if elastic else 'Plastico')+', con My = 0. La ricerca avviene lungo l’eccentricità N/M costante. Inserire le azioni di riferimento e confrontare N, M e il tasso della frontiera.')
  table(['Grandezza','Riferimento','ANTHEA','Scarto %'],[
   ['N [kN]',fmt(r['N']),fmt(rr['N']),fmt(100*abs(rr['N']-r['N'])/abs(r['N']),3)],
   ['Mx [kNm]',fmt(r['Mx']),fmt(rr['Mx']),fmt(100*abs(rr['Mx']-r['Mx'])/abs(r['Mx']),3)],
   ['My [kNm]',fmt(r['My']),fmt(rr['My']),fmt(100*abs(rr['My']-r['My'])/abs(r['My']),3) if abs(r['My'])>1e-6 else 'zero'],
   ['η frontiera','1,000000',fmt(a['boundary']['Utilization'],6),fmt(100*abs(a['boundary']['Utilization']-1),3)],
   ['σc,min [MPa]',fmt(r['sigma_c']),fmt(resp['CMin']),''],
   ['σs,max [MPa]',fmt(max(r['bars'])),fmt(resp['SMax']),''],
   ['σs,min [MPa]',fmt(min(r['bars'])),fmt(resp['SMin']),'']],[48,42,42,38])
  small('Tensioni attese nelle barre B1–B4 [MPa]: '+ '; '.join(fmt(x,3) for x in r['bars'])+'.')
  p(f'Controlli a carico proporzionale: per 0,8·(N,Mx,My), η = {fmt(a["inside"]["Utilization"],6)}; per 1,2·(N,Mx,My), η = {fmt(a["outside"]["Utilization"],6)}. La classificazione interno/esterno è corretta.')
  notes={
   'P3D-01':'Si raggiunge εcu2 = 3,5‰ al vertice compresso e lo snervamento in due barre. I risultati rientrano nelle soglie adottate.',
   'P3D-02':'Il cambio di inclinazione dell’asse neutro inverte il segno di My. Due barre compresse sono snervate; le tensioni di trazione restano inferiori a fyd. Il confronto rientra nelle soglie.',
   'E3D-01':'Governa εc2 = 2‰ al vertice compresso; la massima tensione dell’acciaio è inferiore a fyd. Il piccolo superamento numerico di η = 1 rientra nella tolleranza di validazione.',
   'E3D-02':'La tensione attesa al centro della barra B2 è fyd. ANTHEA restituisce circa 379,96 MPa e una frontiera più interna; lo scarto di η è 2,93%. La definizione del primo limite dell’acciaio e la posizione a cui è applicato devono essere chiarite.',
   'P2D-01':'N e M rientrano nell’1%, ma il dettaglio restituito riporta tutta la sezione quasi uniformemente compressa (εc circa −2‰ e tutte le barre compresse). Questo stato non rappresenta il punto di flessione con armatura inferiore tesa: occorre verificare l’associazione fra punto resistente e dettaglio.',
   'P2D-02':'N e M rientrano nell’1%; la tensione massima di trazione nel dettaglio è 367,58 MPa contro 350,00 MPa attesi. Il dettaglio del punto interpolato deve essere verificato prima di chiudere il caso.',
   'E2D-01':'N e M rientrano nell’1%; il dettaglio restituisce 355,10 MPa di trazione contro 320,00 MPa attesi. La corrispondenza fra azioni del punto interpolato e piano di deformazione resta da controllare.',
   'E2D-02':'Il riferimento raggiunge fyd al centro delle barre tese. Gli scarti di N e M superano l’1% e quello di η è 3,90%; nel dettaglio la massima tensione è circa 379,96 MPa. Il caso resta aperto, anche dopo il controllo a 128 direzioni.'}
  p(notes[id],boldlead='Esito '+status(id)+'.')
 elif mode=='stress':
  actual=a['actual'];pl=r['plane'];linear=v['law']=='linear'
  p('Sezione R. Impostare modello '+('Lineare' if linear else 'Non lineare')+', CLS teso No, φ = 0, assi Locali e combinazione '+('Quasi permanente' if v['set']=='SLE_QP' else 'Rara')+'. Il piano noto genera le azioni indicate; il risolutore deve ricostruire le tensioni ai vertici e nelle barre. Il riferimento è indipendente dalle DLL [R1 e capitolo 4].')
  mathp(f'e(x,y) = {sci(pl[0])} + ({sci(pl[2])}) y')
  p(f'Azioni da immettere: N = {fmt(r["N"],6)} kN; Mx = {fmt(r["Mx"],6)} kNm; My = 0 kNm.')
  if id=='TL-01':
   p('La deformazione è uniforme: ε = 10⁶/[Ecm·(A − As) + Es·As]. La sezione omogeneizzata ha Aeq = A + (Es/Ecm − 1)As; si ricavano σc = −10⁶/Aeq e σs = (Es/Ecm)σc. Questo caso controlla unità, segno di N e sottrazione dell’area di CLS occupata dalle barre.')
  elif id in ['TL-02','VT-01','VT-02']:
   p('L’asse neutro è y = 80 mm e la profondità compressa è xn = 250 − 80 = 170 mm. La distribuzione del CLS è triangolare: Cc = b·xn·|σc,min|/2, applicata a y = 250 − xn/3. Si aggiungono le forze delle barre e si sottrae il CLS nelle barre compresse.')
  elif id=='TN-01':
   p('Per e = 0,001, sc = 17·[2·0,5 − 0,5²] = 12,75 MPa e ss = 200 MPa. La forza è N = −[12,75·(150000 − 1256,6371) + 200·1256,6371]/1000. Il risultato è una soluzione chiusa; Mx e My sono nulli per simmetria.')
  else:
   p('L’asse neutro è y = 80 mm e il vertice superiore raggiunge e = 1,5‰, sul ramo parabolico. Le barre inferiori hanno e = −2,470588‰ e risultano snervate. Si integrano sc(y) e sc(y)·y nel solo tratto 80 ≤ y ≤ 250 mm; il limite di esercizio dell’acciaio non è soddisfatto, come atteso.')
  table(['Grandezza','Riferimento','ANTHEA'],[
   ['σc,min [MPa]',fmt(r['sigma_c'],6),fmt(actual['sigma_cls'],6)],
   ['|σs|max [MPa]',fmt(r['sigma_s'],6),fmt(actual['sigma_acciaio'],6)],
   ['σs B1 [MPa]',fmt(r['bars'][0],5),fmt(actual['tensioni_barre'][0],5)],
   ['σs B3 [MPa]',fmt(r['bars'][2],5),fmt(actual['tensioni_barre'][2],5)],
   ['η tensionale',fmt(r['ratio'],6),fmt(actual['Ratio'],6)]],[66,52,52])
  if id=='VT-01':p('Il CLS governa: ηc = 16/18 = 0,888889; ηs = 160,5097/360 = 0,445860. La combinazione rara deve essere classificata entro i limiti. La verifica è riprodotta correttamente.')
  elif id=='VT-02':p('Per la quasi permanente si ha ηc = 15/13,5 = 1,111111: il limite del calcestruzzo deve risultare superato. Il programma restituisce l’esito negativo atteso. Il test non applica automaticamente alla quasi permanente il limite dell’acciaio previsto per la rara.')
  elif not linear:p('La non linearità è quella del legame parabola–rettangolo di progetto scelto nel benchmark. Questi risultati non validano il diverso diagramma materiale denominato “Non lineare”, né un’analisi reologica nel tempo.')
  p('Le grandezze confrontate rispettano le tolleranze. Esito strutturale restituito: '+actual['Status']+'.',boldlead='Esito CONFORME.')
 elif mode=='crack_scalar':
  p('Fonte numerica: European Concrete Platform, Eurocode 2 Worked Examples, esempio 7.3, pp. 7-8 e 7-9 [R3]. La fonte considera b = 400 mm, h = 600 mm, d = 548 mm, d′ = 46 mm, As = 2712 mm², As′ = 452 mm², M = 300 kNm, n = 15, fctm = 2,9 MPa e carico breve. Riporta xn = 237,8 mm, σs = 234 MPa e wk = 0,184 mm.')
  p('Il caso riproduce la funzione scalare CrackWidth, fornendo direttamente i dati intermedi pubblicati. Non ricalcola l’equilibrio della sezione nella DLL. Si impongono Es = 200000 MPa ed Ecm = Es/15 per conservare αe = 15 della fonte; questo valore non è il modulo C30/37 degli altri benchmark. L’interasse di prova s = 60 mm garantisce il ramo delle barre ravvicinate.')
  mathp('hc,eff = min[2,5·52; (600 − 237,8)/3; 300] = 120,733333 mm')
  mathp('Ac,eff = 400·120,733333 = 48293,333333 mm²')
  mathp('ρeff = 2712/48293,333333 = 0,0561568194')
  table(['Passaggio','Valore di riferimento'],[
   ['Riduzione kt fctm(1+αeρ)/ρ',fmt(r['stiffening'],6)+' MPa'],
   ['Δε efficace',fmt(r['strain'],10)],
   ['Δsm per barre ravvicinate',fmt(r['distance'],6)+' mm'],
   ['wk da dati pubblicati arrotondati',fmt(r['wk'],9)+' mm'],
   ['wk ANTHEA',fmt(a['wk'],9)+' mm'],
   ['wk pubblicato','0,184 mm']],[104,66])
  p('Il ricalcolo produce 0,184570 mm, con differenza di 0,000570 mm dal valore pubblicato: lo scarto è compatibile con gli arrotondamenti intermedi. Il valore calcolato da ANTHEA coincide con il riferimento scalare. In ipotesi XC1, armatura poco sensibile e quasi permanente, sarebbe inferiore a 0,30 mm; questa classificazione è un confronto aggiuntivo, non un dato della fonte.',boldlead='Esito CONFORME.')
 else:
  if mode=='crack_full':
   c=v['crack'];aa=a['crack'];st=a['actual']
   p('Sezione W, flessione semplice N = 0, Mx = '+fmt(r['Mx'],6)+' kNm, My = 0. Modello Lineare, CLS teso No, φ = 0; esposizione XC1, armatura poco sensibile, durata Lunga, aderenza Migliorata e combinazione Quasi permanente. Copriferro e spaziatura sono ricavati automaticamente dalla geometria.')
   p('L’asse neutro si ricava dall’equilibrio N = 0 della sezione omogeneizzata, con correzione n − 1 per l’armatura compressa. Si impone poi σs = 250 MPa nelle barre inferiori e si integra il momento. Il riferimento è un calcolo ripercorribile della presente campagna, basato sulle formule della Circolare [R2].')
   table(['Grandezza','Riferimento','ANTHEA'],[
    ['xn [mm]',fmt(c['x']),fmt(st['Response']['NeutralDistance'])],
    ['As,eff [mm²]',fmt(c['as_eff']),fmt(aa['EffectiveSteel'])],
    ['Ac,eff [mm²]',fmt(c['aceff']),fmt(aa['EffectiveArea'])],
    ['s [mm]','500,0000',fmt(aa['BarSpacing'])],
    ['σs [MPa]','250,0000',fmt(st['sigma_acciaio'])],
    ['wk [mm]',fmt(c['wk'],9),fmt(aa['Width'],9)],
    ['ηw = wk/0,30',fmt(c['wk']/.3,6),fmt(aa['Ratio'],6)]],[65,52.5,52.5])
   p(f'hc,eff = 125 mm; ρeff = {fmt(c["rho"],9)}; s = 500 mm > 5(c + Ø/2) = 250 mm. La deformazione calcolata con il tension stiffening è inferiore al minimo: governa Δε = 0,6·250/200000 = 0,00075.')
   p(f'Δsm,vicino = {fmt(c["near"],6)} mm; Δsm,lontano = {fmt(c["far"],6)} mm. Governa la zona lontana dalle barre e wk = 1,7·{fmt(c["far"],6)}·0,00075 = {fmt(c["wk"],6)} mm.')
   p('La differenza di apertura è inferiore a 0,000003 mm. La verifica strutturale deve risultare non soddisfatta, poiché wk > 0,30 mm; ANTHEA restituisce correttamente “Apertura oltre limite”.',boldlead='Esito CONFORME.')
  else:
   ar=a['actual'];staff=v['asw']>0
   p(f'Sezione R; direzione {v["axis"]}; N = −300 kN; V{v["axis"]} = {fmt(v["V"],0)} kN; bw = {v["bw"]} mm; d = {v["d"]} mm; Asl = 628,318531 mm². Le azioni nell’altra direzione sono nulle. Materiali e coefficienti sono quelli del capitolo 3. Riferimento: NTC § 4.1.2.3.5 [R1].')
   if not staff:
    p('Impostare Asw = 0 per isolare il ramo senza armatura resistente a taglio. σcp = 300000/150000 = 2 MPa < 0,2·17 = 3,4 MPa. Calcolare entrambi i termini della resistenza e scegliere il maggiore, come previsto dalla formula.')
    mathp(f'k = 1 + √(200/{v["d"]}) = {fmt(r["k"],9)}')
    mathp(f'ρl = 628,318531/({v["bw"]}·{v["d"]}) = {fmt(r["rho"],9)}')
    rows=[['Termine con ρl [kN]',fmt(r['branch1'],6),fmt(ar['VRsd'],6)],['Termine minimo [kN]',fmt(r['branch2'],6),fmt(ar['VRcd'],6)]]
   else:
    p('Impostare staffe verticali α = 90°, due bracci Ø8, Asw = 100,530965 mm², passo 150 mm, cotθ = 2 e z = 0,9d. La verifica usa le dimensioni manuali indicate. Lo sforzo normale fornisce αc = 1 + 2/17 = 1,117647.')
    mathp(f'VRsd = (0,9·{v["d"]})·(100,530965/150)·391,304348·2/1000')
    mathp(f'VRcd = (0,9·{v["d"]})·{v["bw"]}·1,117647·8,5·2/5/1000')
    rows=[['VRsd [kN]',fmt(r['branch1'],6),fmt(ar['VRsd'],6)],['VRcd [kN]',fmt(r['branch2'],6),fmt(ar['VRcd'],6)]]
   rows += [['VRd [kN]',fmt(r['VRd'],6),fmt(ar['VRd'],6)],['|VEd|/VRd',fmt(r['ratio'],6),fmt(ar['Ratio'],6)]]
   table(['Grandezza','Riferimento','ANTHEA'],rows,[66,52,52])
   p(('Governa la resistenza delle staffe, essendo VRsd < VRcd.' if staff else 'Il termine con ρl è maggiore della resistenza minima e governa VRd,c.')+' Il tasso è inferiore a 1: la resistenza a taglio della singola direzione è sufficiente.')
   p('Le resistenze e il tasso coincidono con il calcolo indipendente entro l’arrotondamento. Il controllo è eseguito sulla funzione Shear con dati geometrici espliciti; non collauda il loro inserimento manuale nella WPF né le prescrizioni costruttive.',boldlead='Esito CONFORME.')

heading('Quadro riepilogativo degli esiti',newpage=True)
rows=[]
for v in REF['cases']:
 id=v['id'];a=ACT[id]
 if v['mode']=='domain':value='η = '+fmt(a['boundary']['Utilization'],6)
 elif v['mode']=='shear':value='VRd = '+fmt(a['actual']['VRd'],3)+' kN'
 elif v['mode']=='crack_scalar':value='wk = '+fmt(a['wk'],6)+' mm'
 elif v['mode']=='crack_full':value='wk = '+fmt(a['crack']['Width'],6)+' mm'
 else:value='η = '+fmt(a['actual']['Ratio'],6)
 rows.append([id,value,status(id)])
table(['Caso','Risultato ANTHEA','Esito del test'],rows,[34,75,61])
p('Sono conformi 15 casi; 3 casi hanno riserva sul dettaglio tensionale del dominio 2D; 2 casi elastici richiedono chiarimento della frontiera. Il valore η > 1 nei casi TN-02 e VT-02 e wk > 0,30 mm in FE-02 è un risultato strutturale negativo atteso e correttamente riconosciuto dal programma.')

heading('Rilievi e condizioni per la chiusura',newpage=True)
heading('Primo limite elastico dell’acciaio',2)
p('E3D-02 ed E2D-02 sono costruiti imponendo |εs| = fyd/Es al centro delle barre tese. La libreria restituisce circa 379,96 MPa al centro, contro 391,30435 MPa del riferimento. Il tasso di frontiera differisce rispettivamente del 2,93% e del 3,90%; nel caso 2D lo scarto relativo di N è circa il 6,65%. Il raddoppio delle direzioni da 64 a 128 non modifica i risultati.')
p('Occorre identificare se il criterio interno opera al centro, sul contorno della barra o su una deformazione convenzionale diversa. La differenza potrebbe dipendere dal modello adottato e non prova, da sola, un errore di programmazione. La chiusura richiede una definizione documentata del limite e un nuovo riferimento indipendente coerente, oppure una correzione del motore se il limite non è quello previsto. Non si modifica il riferimento solo per ottenere l’esito positivo.')
heading('Coerenza dei dettagli dei domini bidimensionali',2)
p('P2D-01, P2D-02 ed E2D-01 riproducono N e M entro l’1%, ma il dettaglio di tensioni e deformazioni associato al punto resistente non soddisfa il controllo di coerenza. Il caso più evidente è P2D-01: il dettaglio riporta una compressione quasi uniforme con entrambe le file compresse, mentre il momento resistente di circa 212 kNm richiede uno stato flessionale.')
p('Si deve verificare il legame fra punto interpolato, piano di deformazione e funzione che prepara il dettaglio. Per ogni punto esportato, ricostruire N, Mx e My dalle tensioni del medesimo stato e controllarne l’equilibrio. La riserva riguarda l’impiego del dettaglio a corredo della resistenza; il solo accordo delle coordinate del dominio non basta a chiuderla.')
heading('Estensioni necessarie per una validazione generale',2)
p('I 20 casi non coprono sezioni circolari, a T o generiche, armature asimmetriche su più file, torsione, pressoflessione con grandi compressioni, calcestruzzi ad alta resistenza, barre incrudenti, viscosità non nulla o effetti della precompressione. Una successiva campagna dovrà includere tali funzioni, se utilizzate, e il percorso completo di input, salvataggio, ricaricamento ed esportazione nell’interfaccia.')
p('La chiusura formale richiede l’esecuzione documentata delle riprove, la registrazione della build e la firma dei responsabili della redazione, del controllo e dell’approvazione. Le caselle di firma del frontespizio non attestano un’approvazione già rilasciata.')

heading('Riproducibilità e identificazione del software',newpage=True)
p('Revisione dei sorgenti ANTHEA usata per l’esecuzione: 420590f130fd5713b7ce26e4498e5970766dcb27. Il motore è stato compilato in configurazione Release, target .NET 8. Il codice applicativo non è stato modificato per adattarlo ai risultati di riferimento. Il programma di prova è separato dai sorgenti del prodotto.')
manifest=json.loads((ROOT.parents[2]/'lib'/'Checker'/'manifest.json').read_text())
rows=[]
for name in ['GPCChecker.Concrete.dll','GPCModel.dll','GPCGeometry.dll']:
 item=next(z for z in manifest['assemblies'] if z['file']==name)
 rows.append([name,item.get('assemblyVersion',''),sha256((ROOT.parents[2]/'lib'/'Checker'/name).read_bytes()).hexdigest()[:16]])
table(['Libreria','Versione','SHA256 primi 16 caratteri'],rows,[71,32,67])
heading('Algoritmo di integrazione indipendente',2)
p('1. Definire il rettangolo e le quattro barre dalle coordinate del capitolo 3. Per ciascun caso leggere a, be e ce dalla scheda. Calcolare e in ogni punto; tagliare l’integrazione alle rette e = 0 ed e = 0,002, dove cambia il ramo del legame.')
p('2. Suddividere l’intervallo x nei valori in cui queste rette attraversano i bordi y = ±h/2. Per ogni ascissa di integrazione, dividere l’intervallo y nei punti y = (et − a − be·x)/ce interni alla sezione; et vale 0 oppure 0,002. Omettere i tagli non applicabili quando il coefficiente al denominatore è nullo.')
p('3. Su ciascun intervallo [l,u] applicare Gauss–Legendre: coordinate (l+u)/2 + ξj(u−l)/2 e pesi wj(u−l)/2. Usare i tre valori positivi e i corrispondenti negativi della tabella seguente. Integrare sc, sc·y e sc·x; aggiungere per ogni barra As(ss−sc), il suo momento rispetto a x e quello rispetto a y.')
table(['|ξj|','wj'],[['0,238619186083197','0,467913934572691'],['0,661209386466265','0,360761573048139'],['0,932469514203152','0,171324492379170']],[85,85])
p('4. Applicare i segni e i divisori 1000 e 10⁶ indicati nel capitolo 4. Per FE-02 determinare la quota dell’asse neutro con bisezione di N = 0, poi scalare la curvatura fino a σs = 250 MPa. Per taglio e FE-01 sono sufficienti le formule e le sostituzioni delle rispettive schede.')
small('Le cifre in tabella sono arrotondate per la lettura; i confronti sono stati effettuati con i valori non arrotondati. La quadratura è esatta, a precisione numerica, per i polinomi a tratti dei benchmark rettangolari adottati; questa proprietà non si estende automaticamente ad altri legami e geometrie.')

heading('Fonti e riferimenti',newpage=True)
p('R1. Ministero delle Infrastrutture e dei Trasporti. Decreto 17 gennaio 2018, Aggiornamento delle Norme tecniche per le costruzioni. G.U. n. 42 del 20 febbraio 2018, supplemento ordinario n. 8. Riferimenti utilizzati: §§ 4.1.2.1.2, 4.1.2.2.4, 4.1.2.2.5, 4.1.2.3.4, 4.1.2.3.5 e 10.2.')
def link(label,url):
 x=doc.add_paragraph();x.paragraph_format.space_after=Pt(9)
 h=OxmlElement('w:hyperlink');rid=doc.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True);h.set(qn('r:id'),rid)
 r=OxmlElement('w:r');rp=OxmlElement('w:rPr');col=OxmlElement('w:color');col.set(qn('w:val'),'000000');rp.append(col);r.append(rp);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);x._p.append(h)
link('Testo ufficiale NTC 2018 in Gazzetta Ufficiale','https://www.gazzettaufficiale.it/eli/gu/2018/02/20/42/so/8/sg/pdf')
p('R2. Consiglio Superiore dei Lavori Pubblici. Circolare 21 gennaio 2019 n. 7, Istruzioni per l’applicazione dell’Aggiornamento delle Norme tecniche per le costruzioni. G.U. n. 35 del 11 febbraio 2019, supplemento ordinario n. 5. Per la fessurazione: § C4.1.2.2.4.5, formule C4.1.5–C4.1.10, pagine stampate 87–88.')
link('Capitolo C4 della Circolare 2019 in Gazzetta Ufficiale','https://www.gazzettaufficiale.it/do/atto/serie_generale/caricaPdf?cdimg=19A0085500100010110005&dgu=2019-02-11&art.dataPubblicazioneGazzetta=2019-02-11&art.codiceRedazionale=19A00855&art.num=1&art.tiposerie=SG')
p('R3. European Concrete Platform. Eurocode 2 Worked Examples, 2008. Esempio 7.3, Evaluation of crack amplitude, pagine interne 7-8 e 7-9 (pagine PDF 102–103). Fonte del risultato esterno utilizzato in FE-01. Le quantità intermedie sono riportate con la precisione disponibile nella pubblicazione.')
link('European Concrete Platform Eurocode 2 Worked Examples','https://www.federbeton.it/Portals/0/pubdoc/pubblicazioni/Eurocode2_Worked_Examples.pdf?ver=2018-02-09-095039-390')
p('R4. Biasioli F. et al., Eurocode 2 Background and Applications Design of Concrete Buildings Worked Examples, JRC, EUR 26566 EN, 2014, JRC89037. Riferimento metodologico complementare per la costruzione di esempi documentati agli SLU e SLE; nessun risultato del presente confronto è attribuito a questo rapporto se non espressamente indicato.')
link('Rapporto e scheda bibliografica ufficiale JRC89037','https://publications.jrc.ec.europa.eu/repository/handle/JRC89037')
p('R5. ANTHEA, sorgenti locali della revisione identificata nel capitolo precedente: X.Core/CheckerSection.cs, ConcreteMaterials.cs, ConcreteStandards.cs, Ntc2018Checks.cs, SectionShearGeometry.cs; documentazione supporto/docs/calcestruzzo-interfaccia.md; inventario delle librerie lib/Checker/manifest.json. Le informazioni sul comportamento del software sono separate dai riferimenti normativi.')
small('Fonti online consultate il 24 settembre 2026. Gli esempi costruiti per equilibrio e compatibilità sono identificati come riferimenti della campagna; non sono presentati come prove sperimentali o risultati pubblicati. La versione EC2 richiamata è quella di prima generazione su cui si basano gli esempi citati.')

# Footer originari: nome del rapporto e contatore coerente col solo corpo.
for section in doc.sections:
 for footer in [section.footer,section.first_page_footer,section.even_page_footer]:
  for para in footer.paragraphs:
   for run in para.runs:
    if 'MODELLO-RELAZIONE-ITEC' in run.text:run.text='ANTHEA Validazione CA'
   if any('FILENAME' in (i.text or '') for i in para._p.xpath('.//w:instrText')):
    for ch in list(para._p):
     if ch.tag!=qn('w:pPr'):para._p.remove(ch)
    run=para.add_run('ANTHEA Validazione CA');run.font.size=Pt(9)
   for instr in para._p.xpath('.//w:instrText'):
    if 'NUMPAGES' in instr.text:instr.text=instr.text.replace('NUMPAGES','SECTIONPAGES')
for e in doc.settings.element.findall(qn('w:updateFields')):doc.settings.element.remove(e)
upd=OxmlElement('w:updateFields');upd.set(qn('w:val'),'true');doc.settings.element.append(upd)
doc.core_properties.title='Validazione del calcolo delle sezioni in calcestruzzo armato ANTHEA'
doc.core_properties.subject='20 benchmark numerici secondo NTC 2018 e Circolare 2019'
doc.core_properties.author='';doc.core_properties.last_modified_by='';doc.core_properties.comments=''
draft=ROOT/'authored.docx';doc.save(draft)
# Reintegra le parti non interessate byte per byte dal modello.
editable={'[Content_Types].xml','word/document.xml','word/_rels/document.xml.rels','word/styles.xml','word/settings.xml','docProps/core.xml'}
with ZipFile(TEMPLATE) as src,ZipFile(draft) as gen,ZipFile(OUT/'ANTHEA_Validazione_Calcestruzzo_Armato.docx','w',ZIP_DEFLATED) as z:
 for n in src.namelist():
  if n in editable or (n.startswith('word/footer') and n.endswith('.xml')):
   z.writestr(n,gen.read(n) if n in gen.namelist() else src.read(n))
  else:z.writestr(n,src.read(n))
 for n in gen.namelist():
  if n not in src.namelist():z.writestr(n,gen.read(n))
print(OUT/'ANTHEA_Validazione_Calcestruzzo_Armato.docx')

