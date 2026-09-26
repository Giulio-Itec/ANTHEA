from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
from hashlib import sha256
import json, math
from lxml import etree
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

REPO=Path(__file__).resolve().parents[3]
ROOT=REPO/'supporto/artefatti/validazione_ca_2026_09_25'
OUT=REPO/'supporto/documentazione/Validazione_CA_ANTHEA'
OUT.mkdir(exist_ok=True,parents=True)
TEMPLATE=Path(r'C:/Users/g.pacini/Desktop/MODELLO-RELAZIONE-ITEC-AA.docx')
REF=json.loads((ROOT/'reference.json').read_text(encoding='utf8'))
ACT={v['id']:v for v in json.loads((ROOT/'actual_base.json').read_text(encoding='utf8'))}
doc=Document(TEMPLATE)
sections=[deepcopy(s._sectPr) for s in doc.sections]

def fmt(v,n=4):
 if v is None:return '—'
 if abs(v)<1e-9:v=0
 return f'{v:.{n}f}'.replace('.',',')
def sci(v,digits=12):
 if abs(v)<1e-20:return '0'
 s=f'{v:.{digits}g}'
 if 'e' in s:
  mant,exp=s.split('e');s=mant+'×10'+str(int(exp)).translate(str.maketrans('-0123456789','⁻⁰¹²³⁴⁵⁶⁷⁸⁹'))
 return s.replace('.',',')
def status(id):
 if id in ['E3D-02','E2D-02']:return 'DA CHIARIRE'
 if id in ['P2D-01','P2D-02','E2D-01']:return 'CON RISERVA'
 return 'CONFORME'

# Mantiene il frontespizio ITEC e sostituisce i soli slot dimostrativi.
def replace_p(p,text):
 runs=p.runs
 if runs:
  runs[0].text=text
  for r in runs[1:]:r.text=''
 else:p.add_run(text)
replace_p(doc.paragraphs[2],'Validazione del calcolo delle sezioni in calcestruzzo armato')
replace_p(doc.paragraphs[4],'Software ANTHEA  |  NTC 2018 e Circolare 2019')
doc.tables[0].cell(3,0).text='Ambito:'
doc.tables[0].cell(5,1).text='Validazione del software ANTHEA'
for row in doc.tables[1].rows[1:]:
 for c in row.cells:c.text=''
doc.tables[1].cell(0,1).text='Responsabili'
doc.tables[1].cell(1,1).text='Red.'
doc.tables[1].cell(2,1).text='Contr.'
doc.tables[1].cell(3,1).text='Appr.'
for ir in range(1,4):doc.tables[1].cell(ir,2).text='Da indicare'
doc.tables[2].cell(1,0).text='Relazione di validazione numerica'
doc.tables[2].cell(2,0).text='88 esempi numerici\nResistenza e stati limite di esercizio\nTaglio torsione e momento curvatura\nDurabilità ancoraggi e dettagli costruttivi'
for i,vals in enumerate([
 ['Tipo documento: VALIDAZIONE','Nome file:','ANTHEA-VAL-CA-01'],
 ['Software: ANTHEA','Elaborato:','VAL-CA-01']]):
 for j,t in enumerate(vals):doc.tables[3].cell(i,j).text=t
for j,t in enumerate(['01','REVISIONE COMPLETA','25/09/2026','','','']):doc.tables[4].cell(1,j).text=t
for ti,size in [(1,8),(2,11),(3,10),(4,8)]:
 for row in doc.tables[ti].rows:
  for cell in row.cells:
   for para in cell.paragraphs:
    para.alignment=WD_ALIGN_PARAGRAPH.LEFT
    for run in para.runs:run.font.size=Pt(size)

body=doc._element.body
for child in list(body)[23:]:body.remove(child)
group=doc.tables[1]._tbl
group.getparent().remove(group)
# Il modello termina il frontespizio con una sezione continua; l'indice usa il suo pattern corrente.
end=deepcopy(sections[4])
body.append(end)
for s in doc.styles:
 if s.type==WD_STYLE_TYPE.PARAGRAPH and s.name in ['Heading 1','Heading 2','Heading 3']:
  s.font.color.rgb=RGBColor(0,0,0);s.paragraph_format.keep_with_next=True
if 'Title' not in doc.styles:doc.styles.add_style('Title',WD_STYLE_TYPE.PARAGRAPH)
doc.styles['Title'].base_style=doc.styles['Normal'];doc.styles['Title'].font.name='Manrope ExtraBold';doc.styles['Title'].font.size=Pt(18);doc.styles['Title'].font.color.rgb=RGBColor(0,0,0)
doc.paragraphs[2].style=doc.styles['Title']

def p(text='',boldlead=None):
 x=doc.add_paragraph(style='Normal');x.paragraph_format.space_after=Pt(5)
 if boldlead:x.add_run(boldlead+' ').bold=True
 x.add_run(text);return x
def heading(text,level=1,newpage=False):
 x=doc.add_paragraph(text,style=f'Heading {level}')
 if newpage:x.paragraph_format.page_break_before=True
 x.paragraph_format.space_after=Pt(8);return x
def small(text):
 x=p(text)
 for r in x.runs:r.font.size=Pt(10)
 return x
def mathp(text):
 x=doc.add_paragraph();x.paragraph_format.space_after=Pt(6)
 om=OxmlElement('m:oMath');r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=text;r.append(t);om.append(r);x._p.append(om)
 return x
def table(headers,rows,widths=None):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 if widths is None:widths=[170/len(headers)]*len(headers)
 for c,w in zip(t.columns,widths):c.width=Mm(w)
 for i,h in enumerate(headers):t.rows[0].cells[i].text=h
 for row in rows:
  cells=t.add_row().cells
  for i,v in enumerate(row):cells[i].text=str(v)
 for ir,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
  if ir==0:trpr.append(OxmlElement('w:tblHeader'))
  for j,cell in enumerate(row.cells):
   cell.width=Mm(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcpr=cell._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
   for side in ['top','bottom','left','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),'65');e.set(qn('w:type'),'dxa');mar.append(e)
   tcpr.append(mar)
   if ir==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7E6E6');tcpr.append(sh)
   for para in cell.paragraphs:
    para.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_after=Pt(1);para.paragraph_format.space_before=Pt(1)
    for run in para.runs:run.font.size=Pt(10.5);run.bold=ir==0
 bord=OxmlElement('w:tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:
  e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');bord.append(e)
 t._tbl.tblPr.append(bord)
 doc.add_paragraph().paragraph_format.space_after=Pt(0)
 return t

# Indice automatico, solo capitoli per mantenere leggibilità su una pagina.
x=p('INDICE');x.alignment=WD_ALIGN_PARAGRAPH.CENTER;x.runs[0].bold=True
toc=p();f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'TOC \\o "1-1" \\h \\z \\u');toc._p.append(f)
p('I casi sono identificati da un codice univoco riportato nelle schede e nel quadro riepilogativo. La matrice di copertura associa a ogni verifica almeno due esempi; i casi con più controlli espongono separatamente ogni confronto.')
# Chiusura dell'indice e avvio corpo con il pattern ITEC.
idxsect=deepcopy(sections[3]);idxsect.find(qn('w:type')).set(qn('w:val'),'continuous')
idxsect.find(qn('w:pgMar')).set(qn('w:bottom'),'1134')
tocend=p();tocend._p.get_or_add_pPr().append(idxsect)
finalsect=doc.sections[-1]._sectPr
for e in list(finalsect):finalsect.remove(e)
for e in sections[4]:finalsect.append(deepcopy(e))
mar=finalsect.find(qn('w:pgMar'));mar.set(qn('w:bottom'),'1134')
for tag in ['titlePg','pgNumType','type']:
 for e in finalsect.findall(qn('w:'+tag)):finalsect.remove(e)
num=OxmlElement('w:pgNumType');num.set(qn('w:start'),'1');finalsect.append(num)
# Usa il logo dell'intestazione corpo anche nella prima pagina della relazione.
for e in list(finalsect.findall(qn('w:headerReference'))):finalsect.remove(e)
for e in sections[3].findall(qn('w:headerReference')):
 if e.get(qn('w:type'))=='default':finalsect.insert(0,deepcopy(e))

import sys
sys.path.insert(0,str(REPO/'supporto/test/ValidazioneCA20260925'))
import reference_base as rb
EX=json.loads((ROOT/'extra_reference.json').read_text(encoding='utf8'))['cases']
AX={v['id']:v for v in json.loads((ROOT/'actual_extra.json').read_text(encoding='utf8'))}
coverage=[]
def step(text):return small(text)
def compare(rows):
 table(['Grandezza','Riferimento','ANTHEA','Scarto'],[[k,fmt(x,6),fmt(y,6),fmt(y-x,6) if y is not None else '—'] for k,x,y in rows],[53,39,39,39])
def src(text):small('Riferimento: '+text+'. Le fonti R1–R8 sono identificate e collegate nella bibliografia.')
def casehead(v):
 heading(v['id'].replace('-',' ')+' '+v.get('title',v['family']),2)
 src(v.get('ref','R1 §§4.1.2.1.2, 4.1.2.3.4; R2 §C4.1.2.3.4'))
def balance(s,plane,law='parabola'):
 a,b,c=plane
 concrete=rb.response(dict(s,bars=[]),plane,law)
 forces=[]
 for i,(x,y,d) in enumerate(s['bars']):
  e=a+b*x+c*y
  sc=rb.EC*max(e,0) if law=='linear' else rb.FCD*(2*min(1,max(0,e/.002))-min(1,max(0,e/.002))**2)
  ss=rb.ES*e if law=='linear' else max(-rb.FYD,min(rb.FYD,rb.ES*e))
  As=math.pi*d*d/4;force=As*(ss-sc)/1000;forces.append(force)
  if i==0:rows=[]
  rows.append([f'B{i+1}',f'{x:g}; {y:g}',fmt(e*1000,5),fmt(ss,4),fmt(sc,4),fmt(force,5)])
 table(['Barra','x e y mm','e ‰','ss MPa','sc MPa','Fi kN'],rows,[16,33,29,30,30,32])
 step('Per ogni barra As = πØ²/4; Fi = As(ss − sc)/1000. Le colonne ss e sc usano compressione positiva per l’integrazione; nelle viste ANTHEA le tensioni compresse sono negative.')
 cc=-concrete['N'];mx=concrete['Mx'];my=concrete['My']
 step(f'Integrazione del solo CLS: Cc = ∫sc dA/1000 = {fmt(cc,6)} kN; Mcx = ∫sc y dA/10⁶ = {fmt(mx,6)} kNm; Mcy = −∫sc x dA/10⁶ = {fmt(my,6)} kNm.')
 step(f'N = −[Cc + ΣFi] = −[{fmt(cc,6)} + ({fmt(sum(forces),6)})] = {fmt(-cc-sum(forces),6)} kN.')
 step(f'Mx = Mcx + Σ(Fi yi)/1000 = {fmt(mx,6)} + {fmt(sum(f*y/1000 for f,(_,y,_) in zip(forces,s["bars"])),6)} = {fmt(rb.response(s,plane,law)["Mx"],6)} kNm.')
 step(f'My = Mcy − Σ(Fi xi)/1000 = {fmt(my,6)} − ({fmt(sum(f*x/1000 for f,(x,_,_) in zip(forces,s["bars"])),6)}) = {fmt(rb.response(s,plane,law)["My"],6)} kNm.')
def verdict(rows,tol=.01):
 failed=[k for k,x,y in rows if y is None or abs(y-x)>max(1e-5,abs(x)*tol)]
 p('Confronto numerico: '+('entro la tolleranza assegnata.' if not failed else 'scostamento oltre la tolleranza per '+', '.join(failed)+'. I valori sono mantenuti nel rapporto per il controllo del verificatore.'))
 return failed
heading('Oggetto e modalità di utilizzo')
p('La relazione documenta la validazione numerica del modulo ANTHEA per sezioni in calcestruzzo armato, secondo NTC 2018 e Circolare 2019, con integrazioni dell’Eurocodice 2 di prima generazione. Sono sviluppati 88 esempi numerici, oltre a due scenari di controllo dei riscontri manuali e delle rappresentazioni. Ogni caso espone dati, ipotesi, formule, sostituzioni e risultati attesi; le colonne ANTHEA provengono dall’esecuzione delle API correnti del modulo.')
p('Il documento permette al verificatore di ripercorrere i calcoli senza assumere come riferimento il solo risultato del software. Gli esempi sono calcoli analitici o integrazioni numeriche indipendenti; FE 01 riproduce un valore pubblicato. Le fonti normative sono indicate accanto alle formule e raccolte in bibliografia con link e paragrafi.')
p('Un risultato strutturale “non soddisfatto” può essere il risultato corretto di una prova software. La colonna esito strutturale riguarda la sezione del caso; il confronto numerico riguarda la corrispondenza tra riferimento e programma. Gli scostamenti effettivamente rilevati sono riportati, senza attribuire a un test positivo il significato di certificazione di tutte le possibili strutture.')
heading('Attività del verificatore',2)
p('Fissare la build e le librerie; caricare i materiali e le coordinate delle barre indicati; scegliere modello, normativa, assi e combinazione; inserire le azioni; leggere resistenze e dettagli tensionali; confrontare i valori non arrotondati; registrare eventuali differenze. Conservare il foglio e il report con lo stesso identificativo del caso. Ripetere la campagna dopo una variazione del motore o dei criteri di verifica.')
p('Per i domini controllare frontiera, punto interno e punto esterno, e verificare separatamente il piano di deformazione associato al punto. Per momento curvatura controllare N lungo il percorso e la distinzione fra primo snervamento e stato ultimo. Per i controlli costruttivi ripetere entrambe le configurazioni, anche quando una è deliberatamente insufficiente.')
heading('Limiti del documento',2)
p('L’ambito è il modulo di sezione in c.a., non gli altri moduli di ANTHEA né l’analisi globale della struttura. Restano esclusi progetto automatico, sezioni poligonali generiche, fessurazione da analisi non lineare, punzonamento, instabilità, fatica, incendio e verifiche sismiche complete. La sigla interna SLV usata per il dominio elastico non equivale alla verifica sismica dell’opera.')
p('Le condizioni esecutive non desumibili dalla sezione sono riscontri manuali: confermarle non esegue un calcolo aggiuntivo. La fessurazione delle superfici interne dei fori e la ripartizione delle armature secondarie sulle singole facce non sono dichiarate automaticamente verificate.')
heading('Criteri numerici e identificazione')
table(['Oggetto','Tolleranza della campagna'],[['Forze momenti tensioni e curvature','1% del riferimento; presso zero 0,05 kN, 0,01 kNm o 0,01 MPa'],['Formule scalari e dettagli costruttivi','10⁻⁶ relativo o 10⁻⁵ assoluto; confronto prima degli arrotondamenti'],['Apertura fessure','max di 1% e 0,001 mm'],['Esito booleano','Corrispondenza esatta; campo incompleto distinto da esito negativo'],['Rappresentazioni e riscontri manuali','Controlli descritti separatamente dalle prove numeriche']],[58,112])
p('Esecuzione del 25 settembre 2026, .NET 8 Release. Revisione Git di base: 7d1fcbf91196e4d7d28b8449e3a248b71a4d51b3. Il repository contiene modifiche locali: gli hash del motore e dei sorgenti registrati negli artefatti identificano la configurazione effettiva meglio del solo commit.')
manifest=json.loads((REPO/'lib/Checker/manifest.json').read_text(encoding='utf8'))
rows=[]
for name in ['GPCChecker.Concrete.dll','GPCModel.dll','GPCGeometry.dll','GPCUtilities.dll','GPCModelData.dll']:
 item=next((z for z in manifest['assemblies'] if z['file']==name),{})
 rows.append([name,item.get('assemblyVersion',''),sha256((REPO/'lib/Checker'/name).read_bytes()).hexdigest()[:16]])
table(['Libreria','Versione','SHA256 prefisso'],rows,[73,30,67])
heading('Matrice di copertura',newpage=True)
groups={}
for v in REF['cases']+EX:groups.setdefault(v['family'],[]).append(v['id'])
groups['Limitazione tensionale SLE']=['VT-01','VT-03']
groups['Limitazione tensionale SLE_QP']=['VT-02','VT-04']
for k in ['Taglio X','Taglio Y']:
 for sub in ['senza staffe','con staffe']:
  ids=[k.replace('Taglio ','T')+('-01' if sub=='senza staffe' else '-02')]+groups[k+' '+sub]
  groups[k+' '+sub]=ids
 groups.pop(k)
table(['Verifica o funzione numerica','Esempi'],[[k,', '.join(ids)] for k,ids in groups.items()],[103,67])
p('Nei casi DT DP DS DW ogni confronto costruttivo ha due istanze, 01 e 02: interferro; copriferro nominale e margine delle barre; As minima e massima su ciascuna faccia; diametri e interassi longitudinali; diametro, densità e passo delle staffe; armatura ortogonale e relativo interasse. ZG 01 e ZG 02 coprono il limite di armatura nella zona di giunzione.')
heading('Dati comuni e formule di equilibrio',newpage=True)
p('Unità: geometria in mm; aree in mm²; tensioni e moduli in MPa; azioni N e V in kN; momenti e torsione in kNm. N negativo indica compressione. Le coordinate delle barre sono riferite al centro della sezione; il momento Mx è associato alla variazione lungo y e My a quella lungo x con il segno indicato dalle formule.')
p(f'Sezione R: b = 300 mm, h = 500 mm; quattro barre Ø20 nei punti B1 (−100;−200), B2 (100;−200), B3 (−100;200), B4 (100;200). A = 300 × 500 = 150000 mm²; As,i = π20²/4 = {fmt(math.pi*100,6)} mm²; As = 4As,i = {fmt(math.pi*400,6)} mm². Sezione W: b = 600 mm, h = 500 mm, barre (±250;±200), ancora 4Ø20.')
p(f'C30/37, B450C idealizzato elastoplastico salvo diversa indicazione. fcd = 0,85 × 30/1,5 = 17 MPa; fyd = 450/1,15 = {fmt(rb.FYD,9)} MPa; Es = 200000 MPa; Ecm = 22000[(30+8)/10]^0,3 = {fmt(rb.EC,6)} MPa; fctm = 0,3 × 30^(2/3) = {fmt(.3*30**(2/3),6)} MPa.')
mathp('e(x,y) = a + bx + cy;   sc = 17[2t − t²],   t = min(1; max(0; e/0,002))')
mathp('ss = max(−fyd; min(fyd; Es e));   Fi = As,i [ss(ei) − sc(ei)]')
mathp('N = −[∫sc dA + ΣFi]/1000; Mx = [∫sc y dA + ΣFi yi]/10⁶')
mathp('My = −[∫sc x dA + ΣFi xi]/10⁶')
p('e è positivo in compressione nelle formule di integrazione; lo stato esposto da ANTHEA usa ε = −e e tensione negativa in compressione. Il termine −As,i sc evita di contare il calcestruzzo sostituito dalle barre. Nell’analisi lineare sc = Ecm max(e;0), ss = Es e senza plateau. Nei casi non fessurati sc = Ecm e anche in trazione.')
heading('Integrazione ripercorribile',2)
p('Sul rettangolo x ∈ [−b/2;b/2], y ∈ [−h/2;h/2], separare le regioni e ≤ 0, 0 < e < 0,002, e ≥ 0,002. Nella regione parabolica porre A = a/0,002, B = b/0,002, C = c/0,002: sc = 17[2A−A² + (2B−2AB)x + (2C−2AC)y − B²x² − 2BCxy − C²y²]. Integrare questo polinomio, sc x e sc y; nel plateau integrare la costante 17.')
p('Per riprodurre i numeri con quadratura: suddividere x alle intersezioni delle rette e = 0 ed e = 0,002 con y = ±h/2; a ciascuna x suddividere y nei livelli (et−a−bx)/c interni alla sezione. Su ogni intervallo [l,u] applicare i punti (l+u)/2 + ξj(u−l)/2 e i pesi wj(u−l)/2. Usare i valori ±ξ della tabella; ripetere per x e y. Il confronto con ordine 10 conferma le risultanti entro 10⁻⁸ per questi polinomi a tratti.')
table(['|ξj|','wj'],[['0,238619186083197','0,467913934572691'],['0,661209386466265','0,360761573048139'],['0,932469514203152','0,171324492379170']],[85,85])
p('Dominio plastico: emax = 0,0035 per i punti scelti, con trazione CLS esclusa. Dominio elastico convenzionale di riferimento: emax ≤ 0,002 e |ei| ≤ fyd/Es nei centri delle barre. Quest’ultima convenzione è dichiarata esplicitamente perché non coincide sempre con la frontiera numerica restituita dalla DLL; i relativi scarti restano leggibili nei casi E e RE.')
src('R1 §§4.1.2.1.2 e 4.1.2.3.4; R3 §§3.1.7 e 6.1; integrazione polinomiale della campagna')
heading('Domini analisi e verifiche SLE')
for v in REF['cases']:
 id=v['id'];r=v['reference'];a=ACT[id];mode=v['mode'];casehead(v)
 if mode in ['domain','stress','crack_full']:
  sec=v['section'];pl=r['plane']
  p(f'Sezione {"W" if sec["b"]==600 else "R"}; modello {"lineare" if v["law"]=="linear" else "parabola rettangolo"}; CLS teso escluso; φ = 0; assi locali. Piano noto e = {sci(pl[0])} + ({sci(pl[1])})x + ({sci(pl[2])})y.')
  if mode=='domain':
   p(f'Dominio {v["dimension"]}D, {"elastico" if v["state"]=="SLV" else "plastico"}; 64 direzioni; strategia iterativa; eccentricità costante. In 2D: N–M, θ = 0°, My = 0.')
  else:p('Impostare combinazione '+v.get('set','SLE_QP')+'. Le azioni ottenute qui sotto sono gli input del risolutore tensionale.')
  balance(sec,pl,v['law'])
  if mode=='domain':
   rr=a['boundary']['Resistance'];resp=a['boundary']['Response'];rows=[('N kN',r['N'],rr['N']),('Mx kNm',r['Mx'],rr['Mx']),('My kNm',r['My'],rr['My']),('η frontiera',1,a['boundary']['Utilization']),('σc min MPa',r['sigma_c'],resp['CMin']),('σs max MPa',max(r['bars']),resp['SMax']),('σs min MPa',min(r['bars']),resp['SMin'])]
   compare(rows);verdict(rows)
   p(f'Controlli interno ed esterno: azioni 0,8 × (N;Mx;My) ⇒ η atteso 0,8, ANTHEA {fmt(a["inside"]["Utilization"],6)}; azioni 1,2 × (N;Mx;My) ⇒ η atteso 1,2, ANTHEA {fmt(a["outside"]["Utilization"],6)}. Gli esiti devono essere rispettivamente interno ed esterno.')
  else:
   ar=a['actual'];rows=[('σc min MPa',r['sigma_c'],ar['sigma_cls']),('|σs|max MPa',r['sigma_s'],ar['sigma_acciaio']),('η tensionale',r.get('ratio',abs(r['sigma_c'])/13.5),ar['Ratio'])]
   compare(rows);verdict(rows)
   p(f'Limiti C30: rara σc ≤ 0,6 × 30 = 18 MPa e |σs| ≤ 0,8 × 450 = 360 MPa; QP σc ≤ 0,45 × 30 = 13,5 MPa. ηrara = max({fmt(abs(r["sigma_c"])/18,6)};{fmt(r["sigma_s"]/360,6)}), ηQP = {fmt(abs(r["sigma_c"])/13.5,6)}. Si applica il solo criterio della combinazione scelta. Esito ANTHEA: {ar["Status"]}.')
  if mode=='crack_full':
   c=v['crack'];cr=a['crack']
   p(f'XC1, armatura poco sensibile, QP, carico lungo, barre nervate. Equilibrio N = 0 risolto per q: q = {fmt(250-c["x"],6)} mm, xn = {fmt(c["x"],6)} mm; d = 450 mm.')
   for text in [f'hc,eff = min(2,5 × 50; (500−{fmt(c["x"],6)})/3; 250) = {fmt(c["hc"],6)} mm; Ac,eff = 600 × hc = {fmt(c["aceff"],6)} mm².',
    f'As,eff = 2π20²/4 = {fmt(c["as_eff"],6)} mm²; ρ = As,eff/Ac,eff = {fmt(c["rho"],9)}; αe = 200000/{fmt(rb.EC,6)} = {fmt(c["alpha"],6)}.',
    f'Riduzione = 0,4 × {fmt(.3*30**(2/3),6)} × (1 + αeρ)/ρ = {fmt(c["stiffening"],6)} MPa; Δε = max[(250−{fmt(c["stiffening"],6)})/200000;0,6×250/200000] = {fmt(c["strain"],9)}.',
    f'Δsm,v = (3,4×40 + 0,8×0,5×0,425×20/ρ)/1,7 = {fmt(c["near"],6)} mm; s = 500 > 5(40+10)=250 mm; Δsm,l = 0,75(500−xn) = {fmt(c["far"],6)} mm.',
    f'wk = 1,7 × max(Δsm,v;Δsm,l) × Δε = {fmt(c["wk"],9)} mm; ηw = wk/0,30 = {fmt(c["wk"]/.3,6)}; ANTHEA wk = {fmt(cr["Width"],9)} mm.']:step(text)
 elif mode=='crack_scalar':
  c=r;p('Dato pubblicato ECP esempio 7.3: b = 400, h = 600, d = 548, xn = 237,8 mm; As = 2712 mm²; σs = 234 MPa; n = 15; fctm = 2,9 MPa; Es = 200000 MPa; Ecm = Es/15; Ø24; c = 40 mm; durata breve; aderenza migliorata. La fonte riporta wk = 0,184 mm. Qui si verifica la funzione di apertura con quei dati, non l’intero equilibrio della sezione della fonte.')
  for text in ['hc,eff = min[2,5(600−548); (600−237,8)/3; 300] = 120,733333 mm; Ac,eff = 400 × 120,733333 = 48293,333333 mm².',
   f'ρ = 2712/48293,333333 = {fmt(c["rho"],10)}; s = 60 ≤ 5(40+24/2)=260 mm.',
   f'Riduzione = 0,6 × 2,9 × (1+15ρ)/ρ = {fmt(c["stiffening"],6)} MPa.',
   f'Δε = max[(234−{fmt(c["stiffening"],6)})/200000;0,6×234/200000] = {fmt(c["strain"],10)}.',
   f'Δsm = (3,4×40 + 0,8×0,5×0,425×24/ρ)/1,7 = {fmt(c["distance"],6)} mm.',
   f'wk = 1,7 × {fmt(c["distance"],6)} × {fmt(c["strain"],10)} = {fmt(c["wk"],9)} mm; ANTHEA = {fmt(a["wk"],9)} mm.',
   f'Scarto rispetto a 0,184 pubblicato = {fmt(c["wk"]-.184,9)} mm, compatibile con i dati intermedi arrotondati.']:step(text)
  src('R5 esempio 7.3 pp. 7 8 e 7 9, pagine PDF 102 e 103; R2 §C4.1.2.2.4.5 eq. C4.1.5–C4.1.10')
 else:
  # Scalar shear: complete substituted formula, shared with the independent extra reference.
  import reference_extra as re
  par=dict(N=-300,V=v['V'],A=150000,bw=v['bw'],d=v['d'],Asl=2*math.pi*100,Asw=v['asw'])
  expected,steps=re.shear(par);p(f'Direzione {v["axis"]}; bw = {v["bw"]} mm, d = {v["d"]} mm, N = −300 kN; A = 150000 mm²; Asl = 628,318531 mm²; Asw = {fmt(v["asw"],6)} mm²; s = 150 mm; VEd = {v["V"]} kN.')
  for text in steps:step(text)
  compare([(k,expected[k],a['actual'][k]) for k in expected]);src('R1 §4.1.2.3.5 eq. 4.1.23–4.1.28')
def detailing_formula(v,r):
 data=v['input'];par=v['params'];name=r['Name'];bars=data['barre_manuali'];b=data['width_mm'];h=data['height_mm'];Ac=b*h
 areas=[math.pi*z['phi']**2/4 for z in bars];As=sum(areas);phimin=min(z['phi'] for z in bars);phimax=max(z['phi'] for z in bars);F=450/1.15
 n=par.get('N',0);pst=par.get('phiSt',8);s=par.get('s',150);present=par.get('stirrups',True);dur=par.get('dur',25);dev=par.get('dev',10);nom=data['cover_mm'];secondary=par.get('secondary',0);sp=par.get('secondarySpacing',150)
 clear=min(math.hypot(x['x']-y['x'],x['y']-y['y'])-(x['phi']+y['phi'])/2 for i,x in enumerate(bars) for y in bars[i+1:])
 # Every benchmark is a regular rectangular perimeter with two horizontal faces.
 dx=max(np.diff(sorted(set(z['x'] for z in bars))));dy=max(z['y'] for z in bars)-min(z['y'] for z in bars);spacing=max(dx,dy)
 if name=='Interferro minimo':return clear,max(20,phimax,par.get('dg',20)+5),f'slibero = min distanza fra centri − (Øi+Øj)/2 = {fmt(clear)}; richiesto max(20;{phimax};{par.get("dg",20)}+5) = {fmt(max(20,phimax,par.get("dg",20)+5))} mm.'
 if name=='Copriferro nominale':return nom,max(10,dur,pst if present else phimax)+dev,f'cnom,req = max(10;{dur};{pst if present else phimax})+{dev} = {max(10,dur,pst if present else phimax)+dev} mm; input {nom} mm.'
 if name=='Margine copriferro barre longitudinali':
  vals=[min(b/2-abs(z['x']),h/2-abs(z['y']))-z['phi']/2-max(10,dur,z['phi'])-dev for z in bars];margin=min(vals)
  return margin,0,f'min[cgeometrico − (max(10;cmin,dur;Ø)+Δcdev)] = 50 − {phimax}/2 − max(10;{dur};{phimax}) − {dev} = {fmt(margin)} mm; richiesto ≥ 0.'
 if name=='Diametro longitudinale':return phimin,12,f'Ømin = {phimin} mm ≥ 12 mm.'
 if name=='Interasse longitudinale':return spacing,300,f'smax = max(Δx;Δy) = max({fmt(dx)};{fmt(dy)}) = {fmt(spacing)} mm ≤ 300 mm.'
 if name=='Armatura longitudinale minima':return As,max(.1*n*1000/F,.003*Ac),f'As,min = max(0,1 × {n} × 1000/{fmt(F)};0,003 × {Ac}) = {fmt(max(.1*n*1000/F,.003*Ac))} mm².'
 if name in ['Armatura longitudinale massima','Armatura verticale massima','Armatura massima nella giunzione']:
  fraction=.08 if 'giunzione' in name else .04
  return As,fraction*Ac,f'As,max = {fraction} × {Ac} = {fmt(fraction*Ac)} mm².'
 if name=='Diametro staffe':return pst if present else 0,max(6,phimax/4),f'Øst = {pst if present else 0}; limite max(6;{phimax}/4) = {fmt(max(6,phimax/4))} mm.'
 if name=='Passo staffe':return s,min(250,12*phimin),f's = {s}; limite min(250;12 × {phimin}) = {min(250,12*phimin)} mm.'
 if name.startswith(('As,min','As,max','Staffe minime','Passo staffe ·')):
  top='superiore' in name;ids=[i for i,z in enumerate(bars) if (z['y']>=0 if top else z['y']<0)];aa=sum(areas[i] for i in ids);yc=sum(areas[i]*bars[i]['y'] for i in ids)/aa;d=yc+h/2 if top else h/2-yc
  if name.startswith('As,min'):
   fct=.3*30**(2/3);coef=max(.26*fct/450,.0013);return aa,coef*b*d,f'd = {fmt(d)}; coefficiente = max(0,26 × {fmt(fct)}/450;0,0013) = {fmt(coef)}; As,min = {fmt(coef)} × {b} × {fmt(d)} = {fmt(coef*b*d)} mm².'
  if name.startswith('As,max'):return aa,.04*Ac,f'As,max della faccia = 0,04 × {Ac} = {fmt(.04*Ac)} mm².'
  if name.startswith('Staffe minime'):
   val=2*math.pi*pst**2/4*1000/s if present else 0;return val,1.5*b,f'Ast/m = 2 × π × {pst}²/4 × 1000/{s} = {fmt(val)} mm²/m; limite 1,5 × {b} = {1.5*b} mm²/m.'
  return s,min(1000/3,.8*d),f's = {s}; limite min(1000/3;0,8 × {fmt(d)}) = {fmt(min(1000/3,.8*d))} mm.'
 if name=='Trattenimento barre compresse':return s,15*phimin,f's = {s} mm ≤ 15 × {phimin} = {15*phimin} mm.'
 if name=='Armatura secondaria':return secondary,.2*As*1000/b,f'As,principale/m = {fmt(As)} × 1000/{b}; As,secondaria,min = 0,2 × {fmt(As)} × 1000/{b} = {fmt(.2*As*1000/b)} mm²/m.'
 if name=='Interasse armatura principale':
  crit=par.get('critical',False);lim=min((2 if crit else 3)*h,250 if crit else 400);return spacing,lim,f'smax = max({fmt(dx)};{fmt(dy)}) = {fmt(spacing)}; limite min({2 if crit else 3} × {h};{250 if crit else 400}) = {lim} mm.'
 if name=='Interasse armatura secondaria':
  crit=par.get('critical',False);lim=min((3 if crit else 3.5)*h,400 if crit else 450);return sp,lim,f'ssecondario = {sp}; limite min({3 if crit else 3.5} × {h};{400 if crit else 450}) = {lim} mm.'
 if name=='Armatura verticale parete':return As,.002*Ac,f'As,v,min = 0,002 × {Ac} = {fmt(.002*Ac)} mm².'
 if name=='Interasse verticale sulla sezione':return spacing,min(3*min(b,h),400),f'smax = max({fmt(dx)};{fmt(dy)}); limite min(3 × {min(b,h)};400) = {min(3*min(b,h),400)} mm.'
 if name=='Armatura orizzontale per metro':return secondary,max(.25*As*1000/max(b,h),.001*min(b,h)*1000),f'As,h,min = max(0,25 × {fmt(As)} × 1000/{max(b,h)};0,001 × {min(b,h)} × 1000) = {fmt(max(.25*As*1000/max(b,h),.001*min(b,h)*1000))} mm²/m.'
 if name=='Interasse orizzontale':return sp,400,f's = {sp} mm ≤ 400 mm.'
 if r['Passed'] is None:return None,None,'Controllo non numerico: '+r['Explanation']+' Esito atteso da completare, senza una conferma manuale.'
 raise ValueError(name)
import numpy as np
heading('Verifiche estese del modulo',newpage=True)
detail_coverage={}
for v in EX:
 casehead(v);id=v['id'];a=AX[id];mode=v['mode'];expected=v['expected']
 for text in v['steps']:step(text)
 if 'error' in a:
  p('ANTHEA non ha restituito una soluzione: '+a['error'].split('\n')[0]);continue
 actual=a['actual']
 if mode=='extra_detail':
  bars=v['input']['barre_manuali'];As=sum(math.pi*b['phi']**2/4 for b in bars)
  p('Area armature As = '+ ' + '.join(f'π×{b["phi"]}²/4' for b in bars)+f' = {fmt(As,6)} mm².')
  table(['Barra','x mm','y mm','Ø mm'],[[f'B{i+1}',fmt(z['x'],1),fmt(z['y'],1),z['phi']] for i,z in enumerate(bars)],[25,50,50,45])
  summary=[]
  for check in actual:
   name=check['Name'];x,l,formula=detailing_formula(v,check);detail_coverage.setdefault(name,[]).append(id)
   step(name+'. '+formula)
   summary.append([name,fmt(x,3)+' / '+fmt(l,3),fmt(check['Actual'],3)+' / '+fmt(check['Limit'],3),'DA COMPLETARE' if check['Passed'] is None else 'SÌ' if check['Passed'] else 'NO'])
   if x is not None and (abs(x-check['Actual'])>1e-5 or abs(l-check['Limit'])>1e-5):step('Differenza da esaminare fra riferimento geometrico e risultato del programma per '+name+'.')
  table(['Controllo','Rif valore / limite','ANTHEA valore / limite','Soddisfatto'],summary,[66,37,42,25])
 elif mode=='extra_quick':
  rows=[]
  for r in actual:rows.append((r['Direction']+' kNm',(1 if '+' in r['Direction'] else -1)*expected['Mx' if 'x' in r['Direction'] else 'My'],r['Moment']))
  compare(rows);verdict(rows)
  # Explicit section integrations for both directions complete the numerical derivation.
  import reference_extra as re
  N=v['params']['N'];elastic=v['params']['elastic']
  for label,sec in [('Mx',rb.R),('My',dict(b=500,h=300,bars=[[y,x,d] for x,y,d in rb.R['bars']]))]:
   rr,q=re.limit_at_N(sec,N,elastic);p('Scomposizione del riferimento '+label+' positivo');balance(sec,rr['plane'])
 elif mode=='extra_curve':
  table(['f','N kN','M kNm','a','k 1/mm','χ 1/m'],[[z['f'],fmt(z['N'],3),fmt(z['M'],4),sci(z['a'],6),sci(z['k'],6),fmt(z['chi'],7)] for z in expected['trace']],[15,25,32,34,34,30])
  p('Risultati ANTHEA lungo il percorso. Ogni riga è una coppia momento curvatura restituita; il punto aggiuntivo proviene dal raffinamento del primo snervamento.')
  table(['M kNm','χ 1/m','εc max ‰','|εs|max ‰','Stato'],[[fmt(z['Moment'],5),fmt(z['Curvature'],7),fmt(z['ConcreteCompressionStrain'],5),fmt(z['SteelStrain'],5),'Ultimo' if z['Limit'] else 'Snervato' if z['Yielded'] else 'Pre snervamento'] for z in actual['Points']],[37,37,32,32,32])
  compare([('MRd kNm',expected['limit'],actual['LimitMoment']),('χu 1/m',expected['trace'][-1]['chi'],actual['UltimateCurvature'])])
  p(f'N al limite ANTHEA = {fmt(actual["LimitAxialKn"],6)} kN; residuo = {fmt(actual["AxialResidualKn"],6)} kN. Primo snervamento riportato χy = {fmt(actual["YieldCurvature"],7)} 1/m.')
  import reference_extra as re
  ey=(450/1.15)/200000
  ky=re.brentq(lambda k:rb.response(rb.R,[200*k-ey,0,k])['N']-v['params']['N'],ey/450,0.0001)
  ay=200*ky-ey; ry=rb.response(rb.R,[ay,0,ky])
  p(f'Primo snervamento indipendente: εyd = (450/1,15)/200000 = {sci(ey)}. La barra inferiore è a y = −200 mm: a − 200k = −εyd, quindi a = 200k − εyd. Sostituendo nell’equilibrio N(a,k) = {v["params"]["N"]} kN e risolvendo per bisezione si ottengono k = {sci(ky)} 1/mm, a = {sci(ay)}, χy = {fmt(1000*ky,7)} 1/m e My = {fmt(ry["Mx"],6)} kNm.')
  balance(rb.R,[ay,0,ky])
  compare([('χy 1/m',1000*ky,actual['YieldCurvature'])])
  p('Il limite geometrico di curvatura non è una verifica di duttilità sismica dell’elemento.')
 elif mode=='extra_cover':
  compare([('cmin,dur mm',expected['dur'],actual['Cover']['Durability']),('cnom mm',expected['nominal'],actual['Cover']['Nominal'])])
 elif mode=='extra_stress':
  if id.startswith('FT'):
   compare([('wk mm',expected['width'],actual['crack']['Width'])])
   table(['Fascia','Ac efficace mm²','As efficace mm²','wk mm'],[[r['Name'],fmt(r['Area'],3),fmt(r['SteelArea'],3),fmt(r['Width'],7)] for r in actual['crack']['Regions']],[45,45,43,37])
  elif id.startswith(('DE','FF')):
   maximum=max(z['Stress'] for z in actual['ConcreteVertices']);compare([('σct max MPa',expected['maxStress'],maximum)])
   p(f'Esito strutturale atteso: {"soddisfatto" if expected["passed"] else "non soddisfatto"}; ANTHEA: {actual["crack"]["Status"]}.')
  else:
   rows=[(k,x,actual[k]) for k,x in expected.items()];compare(rows);verdict(rows)
   p('Esito tensionale ANTHEA: '+actual['Status']+'. Le tensioni delle barre, nell’ordine delle coordinate e poi dell’eventuale trefolo, sono: '+ '; '.join(fmt(z,6) for z in actual['tensioni_barre'])+' MPa.')
 elif mode=='extra_geometry':
  compare([(k,x,actual[k]) for k,x in expected.items()])
  p('Le inerzie analitiche riportate sono riferimenti aggiuntivi per il controllo delle proprietà geometriche. Il confronto eseguito qui riguarda area parametrica, area effettivamente trasmessa al motore, numero di vertici e fori.')
 else:
  rows=[(k,x,actual[k]) for k,x in expected.items() if not isinstance(x,bool)]
  compare(rows);verdict(rows,1e-6)
  if 'Passed' in expected:p('Esito strutturale atteso: '+('soddisfatto' if expected['Passed'] else 'non soddisfatto')+'; ANTHEA: '+('soddisfatto' if actual['Passed'] else 'non soddisfatto')+'.')
heading('Riscontri manuali e viste di verifica',newpage=True)
p('I seguenti controlli non hanno una formula di resistenza propria: verificano la gestione di dati e risultati già calcolati. Due scenari per ogni riga definiscono il comportamento atteso. Le condizioni geometriche del disegno esecutivo restano responsabilità del verificatore. Questi scenari sono istruzioni riproducibili, non nuove prove numeriche eseguite dalla campagna API.')
table(['Controllo','Scenario A','Scenario B'],[
 ['Confinamento ancoraggi','AN 01 con Da verificare: scheda grigia e nessuna attestazione automatica.','AN 02 con Non conforme: riscontro negativo distinto dalla lunghezza.'],
 ['Posizione e sfalsamento giunzioni','SO 01 con Verificato sul disegno: conferma manuale positiva.','SO 02 con Da verificare: riscontro pendente anche se si allunga la barra.'],
 ['Cautele Ø maggiore di 32','Ø20: campo non applicabile e non mostrato.','Ø40: campo richiesto; un esito della lunghezza non conferma le cautele.'],
 ['Appoggi di estremità','DT 01 non confermato: controllo da completare.','DT 02 confermato: attestazione manuale, non calcolo della traslazione.'],
 ['Barre compresse trattenute','DP 01 da confermare: esito pendente.','DP 02 confermato: resta separata la verifica numerica del passo staffe.'],
 ['Armature secondarie sulle facce','DS 01: totale sufficiente; ripartizione sulle due facce ancora da controllare.','DS 02: totale insufficiente; riscontro sulle facce non sana il totale.'],
 ['Pareti collegamenti tra facce','DW 01: distribuzione verticale verificata separatamente dalle legature.','DW 02: conferme sul disegno non modificano interassi o aree insufficienti.'],
 ['Punzonamento bordi e appoggi soletta','DS 01: campo da completare, senza capacità numerica di punzonamento.','DS 02: identico stato, indipendente dall’esito della flessione.'],
 ['Piano e tensioni della sollecitazione','TL 01: piano uniforme e CLS compresso uniforme.','TL 02: asse neutro y=80 mm, CLS nullo nella zona tesa, barre inferiori tese.'],
 ['Piano e tensioni del punto limite','P3D 01: vertice a −3,5‰, due barre snervate.','P3D 02: My cambia segno coerentemente con la pendenza lungo x.'],
 ['Origine della linea nel dominio','N costante: azione (−500;30;20), origine (−500;0;0).','Eccentricità costante: stessa azione, origine (0;0;0).'],
 ['Fessurazione di sezioni forate','G RF 01: foro sottratto da Ac efficace; superficie interna da verificare.','G CF 01: apertura esterna entro limite non produce un sì globale sulla superficie interna.']
],[48,61,61])
heading('Lettura dei risultati e scostamenti',newpage=True)
p('I confronti sono riportati singolarmente nelle schede: non tutti rientrano nelle tolleranze. Nei domini occorre separare le coordinate resistenti dal dettaglio tensionale associato. La campagna mantiene il riferimento indipendente basato sul centro delle barre, senza modificare i valori attesi per farli coincidere con il programma.')
table(['Casi','Evidenza da controllare'],[
 ['MIN 02','Legame incrudente: il riferimento bilineare dichiarato fornisce 400,891977 MPa, mentre il materiale restituisce 401,255043 MPa. Occorre chiarire la costruzione e la discretizzazione del ramo incrudente prima di dichiarare la concordanza.'],
 ['MC4 01 e MC4 02','Campionamento del legame non lineare: piccoli scarti, circa 0,011086 e 0,004601 MPa, superano la tolleranza stretta assegnata alle formule scalari.'],
 ['E3D 02 e E2D 02','La frontiera elastica restituita è più interna rispetto al riferimento al centro delle barre; tasso e coordinate hanno scarti di alcuni punti percentuali.'],
 ['RE 01 e RE 02','Analogo scarto nelle quattro resistenze elastiche a N assegnato; le schede riportano entrambi i valori.'],
 ['P2D 01 P2D 02 E2D 01','Il dettaglio tensionale associato dal dominio 2D non coincide con lo stato del riferimento, anche quando N e M sono prossimi.'],
 ['Cerchi a 32 o 64 lati','Differenza geometrica intenzionale fra area ideale e poligono inscritto; confrontare il motore con il poligono.'],
 ['Interazione taglio torsione','Le prove TV usano un solo taglio non nullo. L’estensione software con somma delle due quote nel CLS e massimo nell’acciaio è una convenzione del modulo e non dimostra da sola una verifica locale completa delle pareti con Vx e Vy simultanei.']
],[54,116])
p('Gli scostamenti richiedono un riesame delle ipotesi e del risultato prima della chiusura del relativo caso di validazione. Nessun valore ANTHEA è stato sostituito con il valore atteso.')
heading('Riproducibilità della campagna',newpage=True)
p('I file di input e le risposte integrali sono conservati in supporto/artefatti/validazione_ca_2026_09_25. I programmi indipendenti e il programma C# di richiamo delle API sono in supporto/test/ValidazioneCA20260925. I sorgenti di generazione del rapporto sono in supporto/script/validazione_ca_2026_09_25.')
p('Ordine di esecuzione: reference_base.py genera i 20 casi originari; reference_extra.py genera le 68 estensioni; ValidazioneCA.csproj legge separatamente reference.json ed extra_reference.json e scrive actual_base.json e actual_extra.json. I risultati delle API sono acquisiti senza arrotondamento. I calcoli numerici delle schede derivano dai file di riferimento, non dall’output del motore.')
p('Comando del programma di prova: dotnet run --project supporto/test/ValidazioneCA20260925 -c Release -- percorso_input.json percorso_output.json. Per una nuova build conservare la precedente cartella di artefatti e produrre una nuova sottocartella, senza sovrascrivere gli esiti storici.')
heading('Copertura dei singoli dettagli',2)
table(['Controllo del modulo','Casi che lo esercitano'],[[k,', '.join(ids)] for k,ids in detail_coverage.items()],[104,66])
assert all(len(ids)>=2 for ids in detail_coverage.values())
heading('Bibliografia e documentazione esterna',newpage=True)
def link(label,url):
 x=doc.add_paragraph();x.paragraph_format.space_after=Pt(8)
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),doc.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True))
 r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);x._p.append(h)
refs=[
('R1','Ministero delle Infrastrutture e dei Trasporti. Decreto 17 gennaio 2018. Aggiornamento delle Norme tecniche per le costruzioni. G.U. n.42 del 20 febbraio 2018, S.O. n.8. Capitolo 4: materiali, SLE, SLU, taglio, torsione, aderenza e dettagli; §7.9.5.2 per il modello esplicito delle pile; §10.2 per l’impiego dei codici di calcolo.','https://www.gazzettaufficiale.it/eli/gu/2018/02/20/42/so/8/sg/pdf'),
('R2','Consiglio Superiore dei Lavori Pubblici. Circolare 21 gennaio 2019 n.7. G.U. n.35 dell’11 febbraio 2019, S.O. n.5. Capitolo C4: §C4.1.2.2.4.5, formule C4.1.5–C4.1.10, per apertura delle fessure; §C4.1.2.3.4.2 e figura C4.1.12 per momento curvatura; §C4.1.6.1.3, tabella C4.1.IV, per copriferri. Nel PDF separato C4: pagine 8–11; la numerazione stampata appartiene alla pubblicazione integrale.','https://www.gazzettaufficiale.it/do/atto/serie_generale/caricaPdf?cdimg=19A0085500100010110005&dgu=2019-02-11&art.dataPubblicazioneGazzetta=2019-02-11&art.codiceRedazionale=19A00855&art.num=1&art.tiposerie=SG'),
('R3','CEN. EN 1992 1 1:2004 con correzioni applicabili. Eurocode 2 Design of concrete structures Part 1 1 General rules and rules for buildings. Riferimenti: §§3.1.7,3.2.7,4.4,6.1–6.3,7.3.4,8.2,8.4,8.7,8.8,9.2,9.3,9.5,9.6. I rinvii sono all’edizione di prima generazione; non alla revisione 2023. Il testo integrale della norma va consultato nell’edizione regolarmente disponibile al progettista.','https://eurocodes.jrc.ec.europa.eu/EN-Eurocodes/eurocode-2-design-concrete-structures'),
('R4','Biasioli F., Curbach M., Feldmann M., Mancini G., Poljanšek M. Eurocode 2 Background and Applications Design of Concrete Buildings Worked Examples. JRC, EUR 26566 EN, 2014, JRC89037. Capitolo 3 per SLU e SLE; capitolo 4 per dettagli, §4.1 e tabelle 4.1.2–4.1.9 per ancoraggi e sovrapposizioni, pagine PDF 115–118. Le tabelle possono includere riduzioni favorevoli non assunte in AN 01–02: non sono confronti diretti a parità di coefficienti.','https://publications.jrc.ec.europa.eu/repository/handle/JRC89037'),
('R5','European Concrete Platform. Eurocode 2 Worked Examples. 2008. Esempio 7.3 Evaluation of crack amplitude, pagine interne 7 8 e 7 9, pagine PDF 102–103. Dato pubblicato wk = 0,184 mm; riproduzione FE 01 con i dati intermedi arrotondati della fonte.','https://www.federbeton.it/Portals/0/pubdoc/pubblicazioni/Eurocode2_Worked_Examples.pdf?ver=2018-02-09-095039-390'),
('R6','European Commission Joint Research Centre. Eurocode 2 Design of Concrete Buildings Worked Examples. Copia PDF del rapporto R4, utile per raggiungere direttamente le pagine delle tabelle di ancoraggio e dei dettagli.','https://eurocodes.jrc.ec.europa.eu/sites/default/files/2022-06/1110_WS_EC2.pdf')]
for code,text,url in refs:
 p(code+'. '+text);link('Apri la fonte '+code,url)
p('R7. Formule analitiche di geometria e integrazione della presente campagna: area e inerzie dei rettangoli e delle corone circolari; area del poligono regolare inscritto; quadratura di Gauss Legendre per polinomi a tratti. Gli esempi costruiti sono riferimenti ripercorribili e non vengono attribuiti a prove sperimentali o a risultati pubblicati.')
p('R8. Documentazione del comportamento implementato: sorgenti ANTHEA X.Core/CheckerSection.cs, ConcreteDetailing.cs, Ntc2018Checks.cs, ConcreteTensionCracking.cs, ConcreteTorsion.cs, MomentCurvature.cs, SectionMomentResistance.cs; X.Materiali/NtcCover.cs e MinimumConcrete.cs; GPC Model ConcreteMaterialEuropeanCommon e SteelMaterial. Sono fonti di implementazione, non autorità normative indipendenti.')
small('Fonti online consultate il 25 settembre 2026. Ogni scheda distingue il risultato pubblicato, quando disponibile, dal riferimento costruito per questa campagna. Le formule del report sono riscritte con i dati degli esempi; i link rinviano ai documenti completi e ai paragrafi identificati.')
# Preserve template furniture, clear demo file names, and request field refresh.
for section in doc.sections:
 for footer in [section.footer,section.first_page_footer,section.even_page_footer]:
  for para in footer.paragraphs:
   if any('FILENAME' in (i.text or '') for i in para._p.xpath('.//w:instrText')):
    for ch in list(para._p):
     if ch.tag!=qn('w:pPr'):para._p.remove(ch)
    para.add_run('ANTHEA Validazione CA').font.size=Pt(9)
   for run in para.runs:
    if 'MODELLO-RELAZIONE-ITEC' in run.text:run.text='ANTHEA Validazione CA'
   for instr in para._p.xpath('.//w:instrText'):
    if 'NUMPAGES' in instr.text:instr.text=instr.text.replace('NUMPAGES','SECTIONPAGES')
for e in doc.settings.element.findall(qn('w:updateFields')):doc.settings.element.remove(e)
upd=OxmlElement('w:updateFields');upd.set(qn('w:val'),'true');doc.settings.element.append(upd)
doc.core_properties.title='Validazione del calcolo delle sezioni in calcestruzzo armato ANTHEA'
doc.core_properties.subject='Revisione 01 con 88 esempi numerici espliciti'
doc.core_properties.author='';doc.core_properties.last_modified_by=''
draft=ROOT/'authored.docx';doc.save(draft)
final=OUT/'ANTHEA_Validazione_Calcestruzzo_Armato_Rev01.docx'
editable={'[Content_Types].xml','word/document.xml','word/_rels/document.xml.rels','word/styles.xml','word/settings.xml','docProps/core.xml'}
with ZipFile(TEMPLATE) as orig,ZipFile(draft) as gen,ZipFile(final,'w',ZIP_DEFLATED) as z:
 for name in orig.namelist():
  z.writestr(name,gen.read(name) if (name in editable or (name.startswith('word/footer') and name.endswith('.xml'))) and name in gen.namelist() else orig.read(name))
 for name in gen.namelist():
  if name not in orig.namelist():z.writestr(name,gen.read(name))
(ROOT/'coverage.json').write_text(json.dumps(dict(groups=groups,details=detail_coverage),ensure_ascii=False,indent=2),encoding='utf8')
print('FINAL',final)
